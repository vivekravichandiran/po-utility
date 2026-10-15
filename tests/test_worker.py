"""Unit tests for notebooks/po_worker.py (Spark is mocked; XLSX files are real)."""
import os

import openpyxl
import pytest

import po_worker as w

HEADER = "object_type,catalog,schema,table,action\n"


class FakeSpark:
    """Records SQL; raises for statements matching any key in `fail_on`."""

    def __init__(self, fail_on=None):
        self.calls = []
        self.fail_on = fail_on or {}

    def sql(self, text):
        self.calls.append(text)
        for needle, exc in self.fail_on.items():
            if needle in text:
                raise exc
        return None


@pytest.fixture
def vol(tmp_path, monkeypatch):
    """A fake Volume root on local disk. /Volumes/... paths are remapped into tmp_path."""
    root_real = tmp_path / "Volumes" / "cat" / "sch" / "vol"
    root_real.mkdir(parents=True)

    real_open, real_isfile, real_exists = open, os.path.isfile, os.path.exists
    real_makedirs, real_isdir = os.makedirs, os.path.isdir
    real_replace, real_remove = os.replace, os.remove
    real_samefile = os.path.samefile

    def m(p):
        # Only remap real paths. Integer file descriptors (used by zipfile/openpyxl) must pass
        # through untouched, otherwise str(fd) becomes a stray file like "15" in the cwd.
        if not isinstance(p, (str, os.PathLike)):
            return p
        p = os.fspath(p)
        return str(tmp_path) + p if p.startswith("/Volumes/") else p

    import builtins
    import shutil

    import io

    monkeypatch.setattr(builtins, "open", lambda p, *a, **k: real_open(m(p), *a, **k))
    monkeypatch.setattr(io, "open", lambda p, *a, **k: real_open(m(p), *a, **k))
    monkeypatch.setattr(os.path, "isfile", lambda p: real_isfile(m(p)))
    monkeypatch.setattr(os.path, "exists", lambda p: real_exists(m(p)))
    monkeypatch.setattr(os.path, "isdir", lambda p: real_isdir(m(p)))
    monkeypatch.setattr(os.path, "samefile", lambda a, b: real_samefile(m(a), m(b)))
    monkeypatch.setattr(os, "makedirs", lambda p, *a, **k: real_makedirs(m(p), *a, **k))
    monkeypatch.setattr(os, "replace", lambda a, b: real_replace(m(a), m(b)))
    monkeypatch.setattr(os, "remove", lambda p: real_remove(m(p)))
    real_copy = shutil.copyfile
    monkeypatch.setattr(shutil, "copyfile", lambda a, b, **k: real_copy(m(a), m(b), **k))
    return "/Volumes/cat/sch/vol", tmp_path


def write_csv(tmp_path, body, name="ops.csv", header=HEADER):
    p = tmp_path / "Volumes" / "cat" / "sch" / "vol" / name
    p.write_text(header + body, encoding="utf-8")
    return "/Volumes/cat/sch/vol/" + name


def real(tmp_path, vpath):
    return str(tmp_path) + vpath


def run(vol, body, spark=None, header=HEADER, overwrite=False, report="rep.xlsx"):
    root, tmp = vol
    csv_path = write_csv(tmp, body, header=header)
    spark = spark or FakeSpark()
    summary = w.run_worker(
        csv_path, root + "/" + report, root, spark.sql, overwrite=overwrite, job_id="7", run_id="9",
        log=lambda *_: None,
    )
    wb = openpyxl.load_workbook(real(tmp, root + "/" + report))
    details = [dict(zip(w.DETAIL_COLUMNS, r)) for r in wb["Details"].iter_rows(min_row=2, values_only=True)]
    return summary, details, spark, wb


# ----------------------------------------------------------------- SQL generation / quoting
def test_quote_identifier_doubles_backticks():
    assert w.quote_identifier("a`b") == "`a``b`"
    assert w.quote_identifier("x") == "`x`"


@pytest.mark.parametrize(
    "ot,action,args,expected",
    [
        ("CATALOG", "DISABLE", ("finance", "", ""), "ALTER CATALOG `finance` DISABLE PREDICTIVE OPTIMIZATION"),
        ("SCHEMA", "ENABLE", ("sales", "staging", ""), "ALTER SCHEMA `sales`.`staging` ENABLE PREDICTIVE OPTIMIZATION"),
        ("TABLE", "DISABLE", ("sales", "curated", "orders"), "ALTER TABLE `sales`.`curated`.`orders` DISABLE PREDICTIVE OPTIMIZATION"),
        ("TABLE", "ENABLE", ("sales", "curated", "customers"), "ALTER TABLE `sales`.`curated`.`customers` ENABLE PREDICTIVE OPTIMIZATION"),
    ],
)
def test_build_sql_all_scopes_and_actions(ot, action, args, expected):
    assert w.build_sql(ot, action, *args) == expected


def test_build_sql_rejects_non_allowlisted_values():
    with pytest.raises(ValueError):
        w.build_sql("VIEW", "ENABLE", "c", "s", "t")
    with pytest.raises(ValueError):
        w.build_sql("TABLE", "INHERIT", "c", "s", "t")


def test_sql_injection_text_is_quoted_not_executed():
    sql = w.build_sql("TABLE", "ENABLE", "c", "s", "t`;DROP TABLE x;--")
    assert sql == "ALTER TABLE `c`.`s`.`t``;DROP TABLE x;--` ENABLE PREDICTIVE OPTIMIZATION"


# ----------------------------------------------------------------- row validation
def test_validate_row_normalizes_case_and_whitespace():
    r = w.validate_row(2, [" table ", " Sales ", " curated", "orders ", " disable "])
    assert r.status == ""
    assert (r.object_type, r.action) == ("TABLE", "DISABLE")
    assert r.generated_sql.startswith("ALTER TABLE `Sales`.`curated`.`orders` DISABLE")


@pytest.mark.parametrize(
    "fields,fragment",
    [
        (["TABLE", "c", "s", "t", "INHERIT"], "action"),
        (["TABLE", "c", "s", "t", "DROP"], "action"),
        (["VIEW", "c", "s", "t", "ENABLE"], "object_type"),
        (["TABLE", "c", "s", "", "ENABLE"], "table is required"),
        (["SCHEMA", "c", "", "", "ENABLE"], "schema is required"),
        (["CATALOG", "", "", "", "ENABLE"], "catalog is required"),
        (["CATALOG", "c", "s", "", "ENABLE"], "schema must be blank"),
        (["SCHEMA", "c", "s", "t", "ENABLE"], "table must be blank"),
        (["TABLE", "c.s", "s", "t", "ENABLE"], "'.'"),
        (["TABLE", "c", "s p", "t", "ENABLE"], "whitespace"),
        (["TABLE", "c", "s", "a/b", "ENABLE"], "'/'"),
        (["TABLE", "c", "s", "t", "ENABLE", "extra"], "Row has 6 fields"),
        (["TABLE", "c"], "Row has 2 fields"),
    ],
)
def test_validate_row_failures(fields, fragment):
    r = w.validate_row(5, fields)
    assert r.status == w.STATUS_VALIDATION_FAILED
    assert fragment in r.error_message
    assert r.generated_sql == ""


# ----------------------------------------------------------------- duplicate / conflict policy
def test_duplicates_execute_once_and_conflicts_execute_none():
    rows = [
        (2, ["TABLE", "c", "s", "t", "ENABLE"]),
        (3, ["table", "C", "S", "T", "enable"]),  # identical (case-insens) duplicate
        (4, ["SCHEMA", "c", "s", "", "ENABLE"]),
        (5, ["SCHEMA", "c", "s", "", "DISABLE"]),  # conflicts with 4
        (6, ["SCHEMA", "c", "s", "", "ENABLE"]),  # also in conflict group
        (7, ["CATALOG", "c", "", "", "ENABLE"]),
    ]
    records, executable = w.plan_records(rows)
    st = {r.source_row_number: r.status for r in records}
    assert [e.source_row_number for e in executable] == [2, 7]
    assert st[3] == w.STATUS_SKIPPED_DUPLICATE
    assert st[4] == st[5] == st[6] == w.STATUS_SKIPPED_CONFLICT
    by = {r.source_row_number: r for r in records}
    assert by[3].references == [2]
    assert by[4].references == [5, 6]


def test_same_names_different_scope_are_not_duplicates():
    records, executable = w.plan_records(
        [(2, ["CATALOG", "c", "", "", "ENABLE"]), (3, ["SCHEMA", "c", "s", "", "ENABLE"])]
    )
    assert len(executable) == 2


# ----------------------------------------------------------------- execution isolation
def test_execution_failure_does_not_stop_later_rows():
    spark = FakeSpark({"`s2`": PermissionError("PERMISSION_DENIED: no MANAGE")})
    records, executable = w.plan_records(
        [
            (2, ["SCHEMA", "c", "s1", "", "ENABLE"]),
            (3, ["SCHEMA", "c", "s2", "", "ENABLE"]),
            (4, ["SCHEMA", "c", "s3", "", "DISABLE"]),
        ]
    )
    w.execute_records(executable, spark.sql, log=lambda *_: None)
    assert [r.status for r in records] == [w.STATUS_SUCCESS, w.STATUS_EXECUTION_FAILED, w.STATUS_SUCCESS]
    assert len(spark.calls) == 3
    assert "PERMISSION_DENIED" in records[1].error_message
    assert records[1].error_class == "PermissionError"


def test_no_automatic_retry():
    spark = FakeSpark({"`s1`": RuntimeError("boom")})
    _, executable = w.plan_records([(2, ["SCHEMA", "c", "s1", "", "ENABLE"])])
    w.execute_records(executable, spark.sql, log=lambda *_: None)
    assert len(spark.calls) == 1


def test_error_class_uses_spark_condition_when_present():
    class E(Exception):
        def getCondition(self):
            return "SCHEMA_NOT_FOUND"

    assert w._error_class_of(E("x")) == "E:SCHEMA_NOT_FOUND"


# ----------------------------------------------------------------- path validation
def test_path_validation(vol):
    root, _ = vol
    assert w.validate_paths(root + "/a.csv", root + "/b.xlsx", root) == (root + "/a.csv", root + "/b.xlsx")
    for bad_csv, bad_rep, frag in [
        (root + "/a.csv", root + "/a.csv", "must end with .xlsx"),
        (root + "/a.csv", root + "/b.txt", "must end with .xlsx"),
        (root + "/a.txt", root + "/b.xlsx", "must end with .csv"),
        ("/Volumes/other/x/y/a.csv", root + "/b.xlsx", "outside volume_root"),
        (root + "/a.csv", "/tmp/b.xlsx", "/Volumes"),
        (root + "/../vol2/a.csv", root + "/b.xlsx", ".."),
        (root + "/a.csv", root + "/sub/../../b.xlsx", ".."),
        (root + "x/a.csv", root + "/b.xlsx", "outside volume_root"),  # prefix trick
        ("", root + "/b.xlsx", "empty"),
    ]:
        with pytest.raises(w.InputError) as e:
            w.validate_paths(bad_csv, bad_rep, root)
        assert frag in str(e.value), (bad_csv, bad_rep, str(e.value))


def test_report_equal_to_csv_rejected(vol):
    root, _ = vol
    with pytest.raises(w.InputError):
        w.validate_paths(root + "/a.xlsx", root + "/a.xlsx", root)


def test_volume_root_must_be_a_volume():
    with pytest.raises(w.InputError):
        w.validate_volume_root("/Volumes/cat/sch")


# ----------------------------------------------------------------- end-to-end (mock spark, real xlsx)
def test_all_scopes_success_and_report_contents(vol):
    body = (
        "CATALOG,finance,,,DISABLE\n"
        "SCHEMA,sales,staging,,ENABLE\n"
        "TABLE,sales,curated,orders,DISABLE\n"
        "TABLE,sales,curated,customers,ENABLE\n"
    )
    summary, details, spark, wb = run(vol, body)
    assert spark.calls == [
        "ALTER CATALOG `finance` DISABLE PREDICTIVE OPTIMIZATION",
        "ALTER SCHEMA `sales`.`staging` ENABLE PREDICTIVE OPTIMIZATION",
        "ALTER TABLE `sales`.`curated`.`orders` DISABLE PREDICTIVE OPTIMIZATION",
        "ALTER TABLE `sales`.`curated`.`customers` ENABLE PREDICTIVE OPTIMIZATION",
    ]
    assert [d["status"] for d in details] == ["SUCCESS"] * 4
    assert [d["source_row_number"] for d in details] == [2, 3, 4, 5]
    assert summary["overall_outcome"] == w.OUTCOME_ALL_SUCCEEDED
    assert summary["successes"] == 4 and summary["total_input_rows"] == 4
    assert wb.sheetnames == ["Summary", "Details"]
    srows = {r[0]: r[1] for r in wb["Summary"].iter_rows(min_row=2, values_only=True)}
    assert srows["job_id"] == "7" and srows["run_id"] == "9"
    assert srows["report_status"] == "WRITTEN"
    assert details[0]["start_timestamp"].endswith("+00:00")


def test_mixed_scenarios_counts_reconcile(vol):
    body = (
        "TABLE,c,s,ok1,ENABLE\n"
        "TABLE,c,s,ok1,ENABLE\n"  # duplicate
        "SCHEMA,c,s,,ENABLE\n"
        "SCHEMA,c,s,,DISABLE\n"  # conflict pair
        "TABLE,c,s,,ENABLE\n"  # missing table
        "BAD,c,s,t,ENABLE\n"  # bad object type
        "TABLE,c,s,denied,ENABLE\n"  # permission failure
        "TABLE,c,s,after,DISABLE\n"  # still attempted
        "TABLE,c,s,t\n"  # malformed row
    )
    spark = FakeSpark({"`denied`": RuntimeError("PERMISSION_DENIED")})
    summary, details, spark, _ = run(vol, body, spark)
    st = [d["status"] for d in details]
    assert st == [
        "SUCCESS", "SKIPPED_DUPLICATE", "SKIPPED_CONFLICT", "SKIPPED_CONFLICT",
        "VALIDATION_FAILED", "VALIDATION_FAILED", "EXECUTION_FAILED", "SUCCESS", "VALIDATION_FAILED",
    ]
    assert summary["total_input_rows"] == 9
    assert summary["successes"] == 2
    assert summary["execution_failures"] == 1
    assert summary["validation_failures"] == 3
    assert summary["duplicates_skipped"] == 1
    assert summary["conflicts_skipped"] == 2
    assert summary["valid_unique_targets_attempted"] == 3
    assert (
        summary["successes"] + summary["execution_failures"] + summary["validation_failures"]
        + summary["duplicates_skipped"] + summary["conflicts_skipped"] == summary["total_input_rows"]
    )
    assert summary["overall_outcome"] == w.OUTCOME_WITH_FAILURES
    # conflicting / invalid / duplicate rows never reached spark
    assert len(spark.calls) == 3
    assert not any("`s`.`ok1`" in c and c != spark.calls[0] for c in spark.calls)


def test_empty_csv_with_header_produces_workbook(vol):
    summary, details, spark, wb = run(vol, "")
    assert spark.calls == []
    assert details == []
    assert summary["overall_outcome"] == w.OUTCOME_NO_ROWS
    assert wb.sheetnames == ["Summary", "Details"]


def test_blank_lines_ignored_but_row_numbers_preserved(vol):
    summary, details, spark, _ = run(vol, "\nTABLE,c,s,t,ENABLE\n\nSCHEMA,c,s,,ENABLE\n")
    assert [d["source_row_number"] for d in details] == [3, 5]


def test_header_in_different_order_and_bom(vol):
    root, tmp = vol
    p = tmp / "Volumes" / "cat" / "sch" / "vol" / "ops.csv"
    p.write_bytes("\ufeffaction,table,schema,catalog,object_type\nENABLE,t,s,c,TABLE\n".encode("utf-8"))
    summary = w.run_worker(root + "/ops.csv", root + "/r.xlsx", root, FakeSpark().sql, log=lambda *_: None)
    assert summary["successes"] == 1


def test_special_identifier_backtick_is_quoted_and_failure_isolated(vol):
    body = "TABLE,c,s,we`ird,ENABLE\nTABLE,c,s,fine,ENABLE\n"
    spark = FakeSpark({"`we``ird`": RuntimeError("TABLE_OR_VIEW_NOT_FOUND")})
    _, details, spark, _ = run(vol, body, spark)
    assert "`we``ird`" in spark.calls[0]
    assert [d["status"] for d in details] == ["EXECUTION_FAILED", "SUCCESS"]


# ----------------------------------------------------------------- bootstrap failures: no DDL
@pytest.mark.parametrize(
    "header",
    [
        "object_type,catalog,schema,table\n",  # missing action
        "object_type,catalog,schema,table,action,extra\n",  # extra column
        "object_type,catalog,schema,table,action,action\n",  # duplicate
        "type,catalog,schema,table,action\n",  # wrong name
        "\n",
    ],
)
def test_bad_header_fails_before_any_ddl(vol, header):
    root, tmp = vol
    spark = FakeSpark()
    write_csv(tmp, "TABLE,c,s,t,ENABLE\n", header=header)
    with pytest.raises(w.InputError):
        w.run_worker(root + "/ops.csv", root + "/r.xlsx", root, spark.sql, log=lambda *_: None)
    assert spark.calls == []
    assert not os.path.exists(root + "/r.xlsx")


def test_missing_csv_fails_without_ddl(vol):
    root, _ = vol
    spark = FakeSpark()
    with pytest.raises(w.InputError, match="not found"):
        w.run_worker(root + "/nope.csv", root + "/r.xlsx", root, spark.sql, log=lambda *_: None)
    assert spark.calls == []


def test_empty_file_fails(vol):
    root, tmp = vol
    (tmp / "Volumes" / "cat" / "sch" / "vol" / "ops.csv").write_text("")
    with pytest.raises(w.InputError, match="empty"):
        w.run_worker(root + "/ops.csv", root + "/r.xlsx", root, FakeSpark().sql, log=lambda *_: None)


def test_non_utf8_csv_fails(vol):
    root, tmp = vol
    (tmp / "Volumes" / "cat" / "sch" / "vol" / "ops.csv").write_bytes(b"\xff\xfe\x00bad")
    spark = FakeSpark()
    with pytest.raises(w.InputError):
        w.run_worker(root + "/ops.csv", root + "/r.xlsx", root, spark.sql, log=lambda *_: None)
    assert spark.calls == []


def test_report_outside_root_rejected_before_ddl(vol):
    root, tmp = vol
    write_csv(tmp, "TABLE,c,s,t,ENABLE\n")
    spark = FakeSpark()
    with pytest.raises(w.InputError):
        w.run_worker(root + "/ops.csv", "/Volumes/cat/sch/other/r.xlsx", root, spark.sql, log=lambda *_: None)
    assert spark.calls == []


def test_existing_report_not_overwritten_unless_allowed(vol):
    root, tmp = vol
    body = "TABLE,c,s,t,ENABLE\n"
    run(vol, body)
    spark = FakeSpark()
    write_csv(tmp, body)
    with pytest.raises(w.InputError, match="already exists"):
        w.run_worker(root + "/ops.csv", root + "/rep.xlsx", root, spark.sql, log=lambda *_: None)
    assert spark.calls == []  # rejected before DDL
    summary, _, spark2, _ = run(vol, body, overwrite=True)
    assert summary["successes"] == 1 and len(spark2.calls) == 1


def test_input_csv_never_overwritten(vol):
    root, tmp = vol
    path = write_csv(tmp, "TABLE,c,s,t,ENABLE\n")
    before = open(path, "rb").read()
    spark = FakeSpark()
    # report_path pointing at the CSV itself (even with overwrite) must be rejected
    with pytest.raises(w.InputError):
        w.run_worker(path, path, root, spark.sql, overwrite=True, log=lambda *_: None)
    assert spark.calls == []
    assert open(path, "rb").read() == before


# ----------------------------------------------------------------- report failures surface (task must fail)
def test_report_write_failure_raises_report_error(vol, monkeypatch):
    root, tmp = vol
    write_csv(tmp, "TABLE,c,s,t,ENABLE\n")
    spark = FakeSpark()

    def boom(*a, **k):
        raise OSError("disk full")

    monkeypatch.setattr(w, "build_workbook", boom)
    with pytest.raises(w.ReportError, match="disk full"):
        w.run_worker(root + "/ops.csv", root + "/r.xlsx", root, spark.sql, log=lambda *_: None)
    assert len(spark.calls) == 1  # DDL did run; failure is surfaced, not hidden
    assert not os.path.exists(root + "/r.xlsx")


def test_unwritable_report_dir_fails_before_ddl(vol, monkeypatch):
    root, tmp = vol
    write_csv(tmp, "TABLE,c,s,t,ENABLE\n")
    spark = FakeSpark()
    real_open = open

    def deny(p, mode="r", *a, **k):
        if "w" in mode and ".po_probe_" in str(p):
            raise PermissionError("read-only volume")
        return real_open(p, mode, *a, **k)

    import builtins

    monkeypatch.setattr(builtins, "open", deny)
    with pytest.raises(w.InputError, match="not writable") as e:
        w.run_worker(root + "/ops.csv", root + "/sub/r.xlsx", root, spark.sql, log=lambda *_: None)
    assert "IP access list" in str(e.value)  # actionable hint for classic-compute network blocks
    assert spark.calls == []


def test_replace_unsupported_falls_back_to_copy(vol, monkeypatch):
    def nope(a, b):
        raise OSError("rename not supported")

    monkeypatch.setattr(os, "replace", nope)
    summary, details, _, _ = run(vol, "TABLE,c,s,t,ENABLE\n")
    assert summary["successes"] == 1 and len(details) == 1
    leftovers = [f for f in os.listdir(vol[1] / "Volumes/cat/sch/vol") if f.endswith(".tmp")]
    assert leftovers == []


# ----------------------------------------------------------------- spreadsheet safety
def test_formula_injection_neutralized(vol):
    spark = FakeSpark({"`=cmd`": RuntimeError("=HYPERLINK(\"http://evil\",\"x\")")})
    body = "TABLE,c,s,=cmd,ENABLE\nTABLE,c,s,@x,ENABLE\n"
    _, details, _, wb = run(vol, body, spark)
    assert details[0]["table"] == "'=cmd"
    assert details[0]["error_message"].startswith("'=")
    assert details[1]["table"] == "'@x"
    for ws in wb:
        for row in ws.iter_rows():
            for cell in row:
                assert cell.data_type != "f", cell.coordinate


def test_long_error_truncated_and_logged_in_full(vol):
    logs = []
    long_msg = "E" * 40000
    root, tmp = vol
    write_csv(tmp, "TABLE,c,s,t,ENABLE\n")
    w.run_worker(root + "/ops.csv", root + "/r.xlsx", root,
                 FakeSpark({"`t`": RuntimeError(long_msg)}).sql, log=logs.append)
    wb = openpyxl.load_workbook(real(tmp, root + "/r.xlsx"))
    msg = list(wb["Details"].iter_rows(min_row=2, values_only=True))[0][12]
    assert len(msg) <= 32767 and "TRUNCATED" in msg
    assert any(long_msg in entry for entry in logs)


def test_illegal_control_chars_do_not_break_report(vol):
    spark = FakeSpark({"`t`": RuntimeError("bad \x00\x07 char")})
    _, details, _, _ = run(vol, "TABLE,c,s,t,ENABLE\n", spark)
    assert details[0]["status"] == "EXECUTION_FAILED"
    assert "\x00" not in details[0]["error_message"]


def test_compute_type_recorded_in_summary_sheet(vol):
    root, tmp = vol
    write_csv(tmp, "TABLE,c,s,t,ENABLE\n")
    summary = w.run_worker(root + "/ops.csv", root + "/r.xlsx", root, FakeSpark().sql,
                           compute_type="classic", log=lambda *_: None)
    assert summary["compute_type"] == "classic"
    wb = openpyxl.load_workbook(real(tmp, root + "/r.xlsx"))
    assert {r[0]: r[1] for r in wb["Summary"].iter_rows(min_row=2, values_only=True)}["compute_type"] == "classic"
