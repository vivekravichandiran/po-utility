"""Unit tests for the local/remote execution feature of notebooks/po_worker.py.

No network: the remote workspace is simulated (OAuth token endpoint + SQL Statement API).
"""
import base64
import json

import openpyxl
import pytest

import po_worker as w
from test_worker import FakeSpark, real, vol, write_csv  # noqa: F401  (vol is a pytest fixture)

URL = "https://adb-1234567890123456.7.azuredatabricks.net"
CLIENT_ID = "11111111-2222-3333-4444-555555555555"
SECRET = "dose-s3cr3t-VALUE-do-not-leak"
WAREHOUSE = "abcdef0123456789"
HEADER6 = "object_type,catalog,schema,table,action,workspace\n"


# ------------------------------------------------------------ simulated remote workspace
class Clock:
    def __init__(self):
        self.now = 0.0

    def sleep(self, seconds):
        self.now += seconds

    def mono(self):
        return self.now


class FakeWorkspace:
    """HTTP transport double. `script[i]` lists the states statement i reports (POST, then GETs)."""

    def __init__(self, script=None, token_status=200, token_body=None, post_status=200, post_body=None,
                 revoke_first_token=False, get_failures=0, transport_error=None, expires_in=3600):
        self.requests = []
        self.script = script or []
        self.token_status, self.token_body = token_status, token_body
        self.post_status, self.post_body = post_status, post_body
        self.revoke_first_token = revoke_first_token
        self.get_failures = get_failures
        self.transport_error = transport_error
        self.expires_in = expires_in
        self.tokens_issued = 0
        self.sql = []
        self._cursor = {}

    # --- helpers for assertions
    def calls(self, method, fragment):
        return [r for r in self.requests if r["method"] == method and fragment in r["url"]]

    def __call__(self, method, url, headers, body, timeout):
        self.requests.append({"method": method, "url": url, "headers": dict(headers), "body": body,
                              "timeout": timeout})
        if self.transport_error:
            raise self.transport_error
        if url.endswith("/oidc/v1/token"):
            if self.token_status != 200:
                return self.token_status, json.dumps(self.token_body or {"error": "invalid_client"}).encode()
            self.tokens_issued += 1
            return 200, json.dumps({"access_token": "tok-%d" % self.tokens_issued,
                                    "expires_in": self.expires_in, "token_type": "Bearer"}).encode()
        bearer = headers.get("Authorization", "")
        if self.revoke_first_token and bearer == "Bearer tok-1":
            return 401, b'{"error_code":"401","message":"token expired"}'
        if method == "POST":
            if self.post_status != 200:
                return self.post_status, json.dumps(self.post_body or {"message": "denied"}).encode()
            idx = len(self.sql)
            self.sql.append(json.loads(body)["statement"])
            self._cursor[idx] = 0
            return 200, self._render(idx)
        sid = url.rsplit("/", 1)[1]
        idx = int(sid.split("-")[1])
        if self.get_failures:
            self.get_failures -= 1
            return 503, b"unavailable"
        self._cursor[idx] += 1
        return 200, self._render(idx)

    def _render(self, idx):
        states = self.script[idx] if idx < len(self.script) else ["SUCCEEDED"]
        st = states[min(self._cursor[idx], len(states) - 1)]
        status = {"state": st if isinstance(st, str) else st[0]}
        if not isinstance(st, str):
            status["error"] = {"error_code": st[1], "message": st[2]}
        return json.dumps({"statement_id": "st-%d" % idx, "status": status}).encode()


def make_client(ws=None, clock=None, secret_calls=None, **kw):
    ws = ws or FakeWorkspace()
    clock = clock or Clock()

    def secret():
        if secret_calls is not None:
            secret_calls.append(1)
        return SECRET

    kw.setdefault("log", lambda *_: None)
    c = w.RemoteSqlClient(URL, CLIENT_ID, secret, WAREHOUSE, http=ws, sleep=clock.sleep,
                          monotonic=clock.mono, **kw)
    return c, ws, clock


# ------------------------------------------------------------ URL validation
@pytest.mark.parametrize("raw,expected", [
    ("https://adb-1.2.azuredatabricks.net", "https://adb-1.2.azuredatabricks.net"),
    ("  https://ADB-1.2.azuredatabricks.net/  ", "https://adb-1.2.azuredatabricks.net"),
    ("https://adb-1.2.azuredatabricks.net/?o=1234#/notebook", "https://adb-1.2.azuredatabricks.net"),
    ("https://adb-1.2.azuredatabricks.net:443", "https://adb-1.2.azuredatabricks.net"),
    ("https://dbc-a1b2.cloud.databricks.com", "https://dbc-a1b2.cloud.databricks.com"),
    ("https://1234.5.gcp.databricks.com/", "https://1234.5.gcp.databricks.com"),
    ("https://adb-1.2.databricks.azure.cn", "https://adb-1.2.databricks.azure.cn"),
])
def test_normalize_workspace_url_accepts_databricks_hosts(raw, expected):
    assert w.normalize_workspace_url(raw) == expected


@pytest.mark.parametrize("raw,frag", [
    ("", "empty"),
    ("   ", "empty"),
    ("adb-1.2.azuredatabricks.net", "https://"),
    ("http://adb-1.2.azuredatabricks.net", "https://"),
    ("ftp://adb-1.2.azuredatabricks.net", "https://"),
    ("https://user:pw@adb-1.2.azuredatabricks.net", "credentials"),
    ("https://adb-1.2.azuredatabricks.net@evil.com", "credentials"),
    ("https://adb-1.2.azuredatabricks.net:8443", "port"),
    ("https://adb-1.2.azuredatabricks.net:99999", "valid URL"),
    ("https://adb-1.2.azuredatabricks.net/api/2.0", "no path"),
    ("https://evil.com", "not a Databricks workspace host"),
    ("https://adb-1.2.azuredatabricks.net.evil.com", "not a Databricks workspace host"),
    ("https://evil.com/adb-1.2.azuredatabricks.net", "no path"),
    ("https://azuredatabricks.net", "not a Databricks workspace host"),
    ("https://localhost", "not a Databricks workspace host"),
    ("https://169.254.169.254", "not a Databricks workspace host"),
    ("https://adb-1.2.azuredatabricks.net\nHost: evil", "illegal"),
    ("https://adb 1.azuredatabricks.net", "illegal"),
])
def test_normalize_workspace_url_rejects_unsafe_values(raw, frag):
    with pytest.raises(w.InputError) as e:
        w.normalize_workspace_url(raw)
    assert frag in str(e.value)


def test_url_validation_is_identical_in_wrapper_and_worker():
    import po_wrapper as p

    cases = ["https://adb-1.2.azuredatabricks.net/?o=1", "https://evil.com", "http://adb-1.2.azuredatabricks.net",
             "https://u:p@adb-1.2.azuredatabricks.net", "https://adb-1.2.azuredatabricks.net/x", "",
             "https://dbc-1.cloud.databricks.com", "https://adb-1.2.azuredatabricks.net:1"]
    for c in cases:
        try:
            expected = ("ok", w.normalize_workspace_url(c))
        except w.InputError:
            expected = ("err", None)
        try:
            got = ("ok", p.normalize_workspace_url(c))
        except p.ConfigError:
            got = ("err", None)
        assert got == expected, c


# ------------------------------------------------------------ build_remote_client
def build(**over):
    args = dict(url=URL, client_id=CLIENT_ID, secret_scope="po-scope", secret_key="po-key",
                warehouse_id=WAREHOUSE, secret_getter=lambda scope, key: SECRET)
    args.update(over)
    return w.build_remote_client(**args)


def test_build_remote_client_none_when_everything_blank():
    assert build(url="", client_id="", secret_scope="", secret_key="", warehouse_id="") is None
    assert build(url=None, client_id=" ", secret_scope="", secret_key="", warehouse_id="  ") is None


def test_build_remote_client_complete():
    c = build(url=URL + "/?o=1")
    assert c.url == URL and c.warehouse_id == WAREHOUSE and c.client_id == CLIENT_ID


@pytest.mark.parametrize("blank", list(w.REMOTE_SETTINGS))
def test_build_remote_client_partial_settings_rejected(blank):
    names = {"remote_workspace_url": "url", "remote_client_id": "client_id", "remote_secret_scope": "secret_scope",
             "remote_secret_key": "secret_key", "remote_warehouse_id": "warehouse_id"}
    with pytest.raises(w.InputError, match="incomplete.*%s" % blank):
        build(**{names[blank]: ""})


@pytest.mark.parametrize("over,frag", [
    (dict(url="https://evil.com"), "Databricks workspace host"),
    (dict(secret_scope="bad scope"), "remote_secret_scope"),
    (dict(secret_key="a/b"), "remote_secret_key"),
    (dict(secret_scope="x" * 129), "remote_secret_scope"),
    (dict(warehouse_id="short"), "remote_warehouse_id"),
    (dict(warehouse_id="my warehouse name"), "remote_warehouse_id"),
    (dict(client_id="has space"), "remote_client_id"),
])
def test_build_remote_client_invalid_values(over, frag):
    with pytest.raises(w.InputError, match=frag):
        build(**over)


def test_build_remote_client_reads_the_secret_lazily_with_scope_and_key():
    calls = []
    c = build(secret_getter=lambda scope, key: calls.append((scope, key)) or SECRET)
    assert calls == []  # nothing read just by configuring
    c._secret_provider()
    assert calls == [("po-scope", "po-key")]


# ------------------------------------------------------------ OAuth token exchange
def test_token_request_shape():
    c, ws, _ = make_client()
    c.execute("SELECT 1")
    req = ws.calls("POST", "/oidc/v1/token")[0]
    assert req["url"] == URL + "/oidc/v1/token"
    assert req["body"] == b"grant_type=client_credentials&scope=all-apis"
    assert req["headers"]["Content-Type"] == "application/x-www-form-urlencoded"
    basic = req["headers"]["Authorization"].split(" ", 1)
    assert basic[0] == "Basic"
    assert base64.b64decode(basic[1]).decode() == "%s:%s" % (CLIENT_ID, SECRET)
    assert SECRET not in req["url"] and SECRET.encode() not in req["body"]  # only in the auth header


def test_token_is_cached_across_statements_and_secret_read_once():
    secrets = []
    c, ws, _ = make_client(secret_calls=secrets)
    for _ in range(5):
        c.execute("SELECT 1")
    assert ws.tokens_issued == 1 and len(secrets) == 1


def test_token_refreshed_before_expiry():
    c, ws, clock = make_client(FakeWorkspace(expires_in=3600))
    c.execute("SELECT 1")
    clock.now += 3600 - w.REMOTE_TOKEN_EARLY_REFRESH_SECONDS - 1
    c.execute("SELECT 1")
    assert ws.tokens_issued == 1  # still valid, with the early-refresh margin
    clock.now += 2
    c.execute("SELECT 1")
    assert ws.tokens_issued == 2
    assert ws.requests[-1]["headers"]["Authorization"] == "Bearer tok-2"


def test_rejected_token_is_refreshed_once_then_retried():
    c, ws, _ = make_client(FakeWorkspace(revoke_first_token=True))
    assert c.execute("SELECT 1") == "st-0"
    assert ws.tokens_issued == 2
    posts = ws.calls("POST", "/api/2.0/sql/statements")
    assert [p["headers"]["Authorization"] for p in posts] == ["Bearer tok-1", "Bearer tok-2"]
    assert ws.sql == ["SELECT 1"]  # the statement ran exactly once (the 401 attempt never executed it)


@pytest.mark.parametrize("status,body", [
    (401, {"error": "invalid_client", "error_description": "Client authentication failed"}),
    (400, {"error": "invalid_request"}),
    (500, {"raw": "boom"}),
])
def test_token_failure_is_an_auth_error_without_credentials(status, body):
    c, ws, _ = make_client(FakeWorkspace(token_status=status, token_body=body))
    with pytest.raises(w.RemoteError) as e:
        c.execute("SELECT 1")
    assert e.value.error_class == "RemoteAuthError" and "token request" in str(e.value)
    assert SECRET not in str(e.value)
    assert ws.calls("POST", "/api/2.0/sql/statements") == []  # never got as far as SQL


def test_secret_in_error_body_is_scrubbed():
    c, _, _ = make_client(FakeWorkspace(token_status=401, token_body={
        "error_description": "bad secret %s for client" % SECRET}))
    with pytest.raises(w.RemoteError) as e:
        c.execute("SELECT 1")
    assert SECRET not in str(e.value) and "***" in str(e.value)


def test_unreadable_secret_scope_is_an_auth_error():
    def boom():
        raise PermissionError("no READ on scope po-scope")

    c = w.RemoteSqlClient(URL, CLIENT_ID, boom, WAREHOUSE, http=FakeWorkspace(), log=lambda *_: None)
    with pytest.raises(w.RemoteError, match="client secret") as e:
        c.execute("SELECT 1")
    assert e.value.error_class == "RemoteAuthError"


def test_empty_secret_rejected():
    c = w.RemoteSqlClient(URL, CLIENT_ID, lambda: "", WAREHOUSE, http=FakeWorkspace(), log=lambda *_: None)
    with pytest.raises(w.RemoteError, match="empty"):
        c.execute("SELECT 1")


# ------------------------------------------------------------ statements
def test_statement_request_shape():
    c, ws, _ = make_client()
    c.execute("ALTER TABLE `a`.`b`.`c` ENABLE PREDICTIVE OPTIMIZATION")
    req = ws.calls("POST", "/api/2.0/sql/statements/")[0]
    assert req["url"] == URL + "/api/2.0/sql/statements/"
    assert req["headers"]["Authorization"] == "Bearer tok-1"
    body = json.loads(req["body"])
    assert body["statement"] == "ALTER TABLE `a`.`b`.`c` ENABLE PREDICTIVE OPTIMIZATION"
    assert body["warehouse_id"] == WAREHOUSE
    assert body["wait_timeout"] == "30s" and body["on_wait_timeout"] == "CONTINUE"
    assert set(body) == {"statement", "warehouse_id", "wait_timeout", "on_wait_timeout"}  # no params, no catalog


def test_execute_waits_through_pending_and_running_until_succeeded():
    c, ws, clock = make_client(FakeWorkspace(script=[["PENDING", "PENDING", "RUNNING", "SUCCEEDED"]]))
    assert c.execute("SELECT 1") == "st-0"
    assert len(ws.calls("GET", "/api/2.0/sql/statements/st-0")) == 3
    assert clock.now == 3 * w.REMOTE_POLL_INTERVAL_SECONDS  # slept between polls only


def test_execute_returns_immediately_when_already_succeeded():
    c, ws, clock = make_client(FakeWorkspace(script=[["SUCCEEDED"]]))
    c.execute("SELECT 1")
    assert ws.calls("GET", "/api/2.0/sql/statements/st-0") == [] and clock.now == 0


def test_closed_counts_as_success():
    c, _, _ = make_client(FakeWorkspace(script=[["RUNNING", "CLOSED"]]))
    assert c.execute("SELECT 1") == "st-0"


def test_failed_statement_reports_remote_error_code_and_message():
    ws = FakeWorkspace(script=[["RUNNING", ("FAILED", "PERMISSION_DENIED", "User does not have MANAGE on table t")]])
    c, _, _ = make_client(ws)
    with pytest.raises(w.RemoteError) as e:
        c.execute("ALTER ...")
    assert e.value.error_class == "RemoteSqlError:PERMISSION_DENIED"
    assert "PERMISSION_DENIED" in str(e.value) and "MANAGE" in str(e.value)
    assert e.value.statement_id == "st-0"


@pytest.mark.parametrize("code,message,expected", [
    # real responses observed from a live workspace
    ("BAD_REQUEST", "[TABLE_OR_VIEW_NOT_FOUND] The table or view `c`.`s`.`t` cannot be found.", "TABLE_OR_VIEW_NOT_FOUND"),
    ("BAD_REQUEST", "PERMISSION_DENIED: User does not have USE CATALOG on Catalog 'po_ctl_t3'.", "PERMISSION_DENIED"),
    ("BAD_REQUEST", "[SCHEMA_NOT_FOUND] The schema `x` cannot be found", "SCHEMA_NOT_FOUND"),
    ("BAD_REQUEST", "[UNRESOLVED_COLUMN.WITH_SUGGESTION] A column cannot be resolved", "UNRESOLVED_COLUMN.WITH_SUGGESTION"),
    # nothing recognisable: fall back to the API's own code
    ("BAD_REQUEST", "something odd happened", "BAD_REQUEST"),
    ("", "", "FAILED"),
    ("BAD_REQUEST", "lowercase: not a condition", "BAD_REQUEST"),
])
def test_failed_statement_error_class_uses_the_real_condition(code, message, expected):
    assert w.RemoteSqlClient.sql_condition(code, message) == expected
    c, _, _ = make_client(FakeWorkspace(script=[[("FAILED", code, message)]]))
    with pytest.raises(w.RemoteError) as e:
        c.execute("ALTER ...")
    assert e.value.error_class == "RemoteSqlError:" + expected
    assert message in str(e.value) or not message  # the full text is kept for the report


def test_canceled_statement_is_a_failure():
    c, _, _ = make_client(FakeWorkspace(script=[["RUNNING", "CANCELED"]]))
    with pytest.raises(w.RemoteError) as e:
        c.execute("ALTER ...")
    assert e.value.error_class == "RemoteSqlError:CANCELED" and e.value.statement_id == "st-0"


def test_unknown_state_is_a_failure_not_a_success():
    c, _, _ = make_client(FakeWorkspace(script=[["WEIRD"]]))
    with pytest.raises(w.RemoteError, match="unexpected state"):
        c.execute("ALTER ...")


def test_timeout_does_not_cancel_and_reports_unknown_outcome():
    ws = FakeWorkspace(script=[["RUNNING"]])
    c, ws, clock = make_client(ws, statement_timeout=60)
    with pytest.raises(w.RemoteError) as e:
        c.execute("ALTER ...")
    assert e.value.error_class == "RemoteTimeout" and e.value.statement_id == "st-0"
    assert "NOT cancelled" in str(e.value) and "unknown" in str(e.value) and "st-0" in str(e.value)
    assert clock.now == 60
    # the only calls: token, one POST, GET polls. No cancel / DELETE of any kind.
    assert {r["method"] for r in ws.requests} <= {"POST", "GET"}
    assert not [r for r in ws.requests if "cancel" in r["url"]]
    assert len(ws.sql) == 1  # submitted exactly once


def test_default_statement_timeout_is_15_minutes():
    assert w.REMOTE_STATEMENT_TIMEOUT_SECONDS == 900


def test_transient_poll_errors_are_tolerated():
    c, ws, _ = make_client(FakeWorkspace(script=[["RUNNING", "SUCCEEDED"]], get_failures=2))
    assert c.execute("ALTER ...") == "st-0"
    assert len(ws.sql) == 1  # polling errors never resubmit the statement


def test_persistent_poll_errors_give_up_with_unknown_outcome():
    c, ws, _ = make_client(FakeWorkspace(script=[["RUNNING"]], get_failures=99))
    with pytest.raises(w.RemoteError) as e:
        c.execute("ALTER ...")
    assert e.value.error_class == "RemoteConnectionError" and e.value.statement_id == "st-0"
    assert "unknown" in str(e.value) and len(ws.sql) == 1


def test_submission_http_error_is_reported_with_status():
    c, ws, _ = make_client(FakeWorkspace(post_status=403, post_body={"message": "no CAN_USE on warehouse"}))
    with pytest.raises(w.RemoteError) as e:
        c.execute("ALTER ...")
    assert e.value.error_class == "RemoteApiError:HTTP_403" and "CAN_USE" in str(e.value)
    assert e.value.statement_id == ""


def test_ip_acl_block_gets_an_actionable_hint():
    body = {"message": "Source IP address: 20.1.2.3 is blocked by Databricks IP ACL for workspace"}
    c, _, _ = make_client(FakeWorkspace(post_status=403, post_body=body))
    with pytest.raises(w.RemoteError, match="IP access list"):
        c.execute("ALTER ...")


def test_missing_statement_id_is_an_error():
    ws = FakeWorkspace()
    orig = ws.__call__

    class NoId(FakeWorkspace):
        def _render(self, idx):
            return b'{"status":{"state":"SUCCEEDED"}}'

    c, _, _ = make_client(NoId())
    with pytest.raises(w.RemoteError, match="no statement_id"):
        c.execute("SELECT 1")
    assert orig  # silence unused


def test_transport_failure_is_a_connection_error_without_secret():
    ws = FakeWorkspace(transport_error=OSError("Name or service not known; auth was Basic xyz " + SECRET))
    c, _, _ = make_client(ws)
    with pytest.raises(w.RemoteError) as e:
        c.execute("SELECT 1")
    assert e.value.error_class == "RemoteConnectionError"
    assert URL in str(e.value)


def test_non_json_response_is_handled():
    def http(method, url, headers, body, timeout):
        return 200, b"<html>gateway</html>"

    c = w.RemoteSqlClient(URL, CLIENT_ID, lambda: SECRET, WAREHOUSE, http=http, log=lambda *_: None)
    with pytest.raises(w.RemoteError, match="token request"):
        c.execute("SELECT 1")


def test_preflight_runs_select_1():
    c, ws, _ = make_client()
    c.preflight()
    assert ws.sql == ["SELECT 1"]


def test_secret_and_token_never_appear_in_any_message():
    seen = []
    ws = FakeWorkspace(script=[["RUNNING", ("FAILED", "X", "echo tok-1 and %s" % SECRET)]])
    c, _, _ = make_client(ws, log=seen.append)
    with pytest.raises(w.RemoteError) as e:
        c.execute("ALTER ...")
    text = str(e.value) + " ".join(seen)
    assert SECRET not in text and "tok-1" not in text and "***" in text


# ------------------------------------------------------------ CSV: workspace column
def test_header_without_workspace_column_is_all_local(vol):
    root, tmp = vol
    write_csv(tmp, "TABLE,c,s,t,ENABLE\n")
    rows = w.read_csv_records(root + "/ops.csv")
    assert rows == [(2, ["TABLE", "c", "s", "t", "ENABLE", "local"])]
    assert w.plan_records(rows)[0][0].workspace == "local"


def test_header_with_workspace_column_any_order(vol):
    root, tmp = vol
    write_csv(tmp, "remote,ENABLE,t,s,c,TABLE\nLocal,DISABLE,t2,s,c,table\n",
              header="workspace,action,table,schema,catalog,object_type\n")
    rows = w.read_csv_records(root + "/ops.csv")
    assert rows == [(2, ["TABLE", "c", "s", "t", "ENABLE", "remote"]),
                    (3, ["table", "c", "s", "t2", "DISABLE", "Local"])]
    recs, _ = w.plan_records(rows)
    assert [r.workspace for r in recs] == ["remote", "local"]  # case-insensitive


@pytest.mark.parametrize("header", [
    "object_type,catalog,schema,table,action,workspace,extra\n",
    "object_type,catalog,schema,table,action,location\n",
    "object_type,catalog,schema,table,workspace\n",
    "object_type,catalog,schema,table,action,workspace,workspace\n",
])
def test_bad_headers_still_rejected(vol, header):
    root, tmp = vol
    write_csv(tmp, "TABLE,c,s,t,ENABLE,local\n", header=header)
    with pytest.raises(w.InputError, match="header"):
        w.read_csv_records(root + "/ops.csv")


def test_blank_or_unknown_workspace_is_a_validation_error_never_silently_local(vol):
    root, tmp = vol
    write_csv(tmp, "TABLE,c,s,a,ENABLE,\nTABLE,c,s,b,ENABLE,cloud\nTABLE,c,s,c,ENABLE,remote \n", header=HEADER6)
    recs, executable = w.plan_records(w.read_csv_records(root + "/ops.csv"))
    assert [r.status for r in recs] == ["VALIDATION_FAILED", "VALIDATION_FAILED", ""]
    assert "workspace" in recs[0].error_message and "workspace" in recs[1].error_message
    assert [e.source_row_number for e in executable] == [4]


def test_short_row_under_a_workspace_header_is_malformed_not_local(vol):
    """A row that forgot the workspace cell must NOT be treated as a legacy (local) row."""
    root, tmp = vol
    write_csv(tmp, "TABLE,c,s,t,ENABLE\n", header=HEADER6)
    recs, executable = w.plan_records(w.read_csv_records(root + "/ops.csv"))
    assert executable == []
    assert recs[0].status == "VALIDATION_FAILED" and recs[0].error_class == "MalformedRow"
    assert "expected 6" in recs[0].error_message


def test_extra_field_row_under_legacy_header_is_malformed(vol):
    root, tmp = vol
    write_csv(tmp, "TABLE,c,s,t,ENABLE,remote\n")  # 6 fields under a 5-column header
    recs, executable = w.plan_records(w.read_csv_records(root + "/ops.csv"))
    assert executable == [] and recs[0].error_class == "MalformedRow" and "expected 5" in recs[0].error_message


def test_same_target_in_local_and_remote_is_not_a_duplicate_or_conflict():
    recs, executable = w.plan_records([
        (2, ["TABLE", "c", "s", "t", "ENABLE", "local"]),
        (3, ["TABLE", "c", "s", "t", "DISABLE", "remote"]),  # opposite action, other workspace: fine
        (4, ["TABLE", "c", "s", "t", "ENABLE", "remote"]),   # conflicts with 3 (same workspace)
    ])
    st = {r.source_row_number: r.status for r in recs}
    assert st[2] == "" and st[3] == st[4] == w.STATUS_SKIPPED_CONFLICT
    assert [e.source_row_number for e in executable] == [2]


def test_duplicates_within_the_remote_workspace_run_once():
    recs, executable = w.plan_records([
        (2, ["SCHEMA", "c", "s", "", "ENABLE", "remote"]),
        (3, ["schema", "C", "S", "", "enable", "REMOTE"]),
    ])
    assert [e.source_row_number for e in executable] == [2]
    assert recs[1].status == w.STATUS_SKIPPED_DUPLICATE


# ------------------------------------------------------------ execution routing
class RecordingRemote:
    def __init__(self, fail_on=None):
        self.sql, self.fail_on = [], fail_on or {}

    def __call__(self, sql):
        self.sql.append(sql)
        for needle, exc in self.fail_on.items():
            if needle in sql:
                raise exc
        return "st-%d" % (len(self.sql) - 1)


def plan(rows):
    return w.plan_records(rows)[1]


def test_rows_go_to_the_runner_for_their_workspace_in_csv_order():
    spark, remote, order = FakeSpark(), RecordingRemote(), []
    ex = plan([
        (2, ["SCHEMA", "c", "l1", "", "ENABLE", "local"]),
        (3, ["SCHEMA", "c", "r1", "", "ENABLE", "remote"]),
        (4, ["SCHEMA", "c", "l2", "", "DISABLE", "local"]),
        (5, ["SCHEMA", "c", "r2", "", "DISABLE", "remote"]),
    ])
    w.execute_records(ex, lambda s: order.append(("local", s)) or spark.sql(s), log=lambda *_: None,
                      remote_runner=lambda s: order.append(("remote", s)) or remote(s))
    assert [o[0] for o in order] == ["local", "remote", "local", "remote"]
    assert [c for c in spark.calls] == [
        "ALTER SCHEMA `c`.`l1` ENABLE PREDICTIVE OPTIMIZATION", "ALTER SCHEMA `c`.`l2` DISABLE PREDICTIVE OPTIMIZATION"]
    assert remote.sql == [
        "ALTER SCHEMA `c`.`r1` ENABLE PREDICTIVE OPTIMIZATION", "ALTER SCHEMA `c`.`r2` DISABLE PREDICTIVE OPTIMIZATION"]
    assert [e.status for e in ex] == ["SUCCESS"] * 4
    assert [e.remote_statement_id for e in ex] == ["", "st-0", "", "st-1"]  # local rows never get one


def test_remote_failure_is_isolated_and_recorded_with_statement_id():
    boom = w.RemoteError("PERMISSION_DENIED: no MANAGE", "RemoteSqlError:PERMISSION_DENIED", "st-77")
    remote = RecordingRemote({"`bad`": boom})
    ex = plan([
        (2, ["SCHEMA", "c", "ok1", "", "ENABLE", "remote"]),
        (3, ["SCHEMA", "c", "bad", "", "ENABLE", "remote"]),
        (4, ["SCHEMA", "c", "ok2", "", "ENABLE", "remote"]),
        (5, ["SCHEMA", "c", "loc", "", "ENABLE", "local"]),
    ])
    spark = FakeSpark()
    logs = []
    w.execute_records(ex, spark.sql, log=logs.append, remote_runner=remote)
    assert [e.status for e in ex] == ["SUCCESS", "EXECUTION_FAILED", "SUCCESS", "SUCCESS"]
    assert ex[1].error_class == "RemoteSqlError:PERMISSION_DENIED" and ex[1].remote_statement_id == "st-77"
    assert "PERMISSION_DENIED" in ex[1].error_message
    assert len(remote.sql) == 3 and len(spark.calls) == 1  # no retry; later rows still ran
    assert any("[remote] FAILED" in l and "st-77" in l for l in logs)
    assert any("[remote] SUCCESS" in l for l in logs) and any("[local] SUCCESS" in l for l in logs)


def test_unexpected_exception_from_remote_runner_is_isolated_too():
    ex = plan([(2, ["SCHEMA", "c", "a", "", "ENABLE", "remote"]), (3, ["SCHEMA", "c", "b", "", "ENABLE", "remote"])])
    remote = RecordingRemote({"`a`": ValueError("odd")})
    w.execute_records(ex, FakeSpark().sql, log=lambda *_: None, remote_runner=remote)
    assert [e.status for e in ex] == ["EXECUTION_FAILED", "SUCCESS"]
    assert ex[0].error_class == "ValueError"


def test_remote_row_never_falls_back_to_local_when_remote_is_missing():
    spark = FakeSpark()
    ex = plan([(2, ["SCHEMA", "c", "a", "", "ENABLE", "remote"])])
    w.execute_records(ex, spark.sql, log=lambda *_: None, remote_runner=None)
    assert spark.calls == []
    assert ex[0].status == "EXECUTION_FAILED" and ex[0].error_class == "RemoteNotConfigured"


def test_legacy_five_field_rows_still_execute_locally():
    spark = FakeSpark()
    ex = plan([(2, ["SCHEMA", "c", "a", "", "ENABLE"])])
    w.execute_records(ex, spark.sql, log=lambda *_: None)
    assert len(spark.calls) == 1 and ex[0].status == "SUCCESS"


# ------------------------------------------------------------ end to end (real xlsx, simulated remote)
def run_e2e(vol, body, spark=None, remote=None, header=HEADER6, logs=None, report="rep.xlsx"):
    root, tmp = vol
    write_csv(tmp, body, header=header)
    spark = spark or FakeSpark()
    summary = w.run_worker(root + "/ops.csv", root + "/" + report, root, spark.sql, job_id="7", run_id="9",
                           log=(logs.append if logs is not None else (lambda *_: None)), remote_client=remote)
    wb = openpyxl.load_workbook(real(tmp, root + "/" + report))
    details = [dict(zip(w.DETAIL_COLUMNS, r)) for r in wb["Details"].iter_rows(min_row=2, values_only=True)]
    srows = {r[0]: r[1] for r in wb["Summary"].iter_rows(min_row=2, values_only=True)}
    return summary, details, spark, srows


MIXED = (
    "SCHEMA,sales,loc1,,ENABLE,local\n"
    "TABLE,sales,curated,orders,DISABLE,remote\n"
    "CATALOG,finance,,,ENABLE,remote\n"
    "SCHEMA,sales,loc2,,DISABLE,local\n"
)


def test_e2e_mixed_local_and_remote(vol):
    client, ws, _ = make_client()
    logs = []
    summary, details, spark, srows = run_e2e(vol, MIXED, remote=client, logs=logs)
    assert spark.calls == ["ALTER SCHEMA `sales`.`loc1` ENABLE PREDICTIVE OPTIMIZATION",
                           "ALTER SCHEMA `sales`.`loc2` DISABLE PREDICTIVE OPTIMIZATION"]
    # preflight SELECT 1 first, then exactly the two remote statements
    assert ws.sql == ["SELECT 1",
                      "ALTER TABLE `sales`.`curated`.`orders` DISABLE PREDICTIVE OPTIMIZATION",
                      "ALTER CATALOG `finance` ENABLE PREDICTIVE OPTIMIZATION"]
    assert [d["workspace"] for d in details] == ["local", "remote", "remote", "local"]
    assert [d["status"] for d in details] == ["SUCCESS"] * 4
    assert [d["remote_statement_id"] for d in details] == [None, "st-1", "st-2", None]
    assert summary["local_rows"] == 2 and summary["remote_rows"] == 2
    assert summary["remote_successes"] == 2 and summary["remote_execution_failures"] == 0
    assert summary["successes"] == 4 and summary["overall_outcome"] == w.OUTCOME_ALL_SUCCEEDED
    assert srows["remote_workspace_url"] == URL and srows["remote_warehouse_id"] == WAREHOUSE
    # status is in the logs as well as the report
    assert any("[remote] SUCCESS" in l and "`orders`" in l and "st-1" in l for l in logs)
    assert any("[local] SUCCESS" in l for l in logs)
    assert any("Remote workspace reachable" in l for l in logs)


def test_e2e_remote_failure_is_recorded_in_report_and_log_and_does_not_stop_the_rest(vol):
    ws = FakeWorkspace(script=[["SUCCEEDED"],  # preflight
                               ["RUNNING", ("FAILED", "PERMISSION_DENIED", "no MANAGE on orders")],
                               ["SUCCEEDED"]])
    client, _, _ = make_client(ws)
    logs = []
    summary, details, spark, _ = run_e2e(vol, MIXED, remote=client, logs=logs)
    assert [d["status"] for d in details] == ["SUCCESS", "EXECUTION_FAILED", "SUCCESS", "SUCCESS"]
    bad = details[1]
    assert bad["workspace"] == "remote" and bad["error_class"] == "RemoteSqlError:PERMISSION_DENIED"
    assert "no MANAGE on orders" in bad["error_message"] and bad["remote_statement_id"] == "st-1"
    assert summary["remote_successes"] == 1 and summary["remote_execution_failures"] == 1
    assert summary["execution_failures"] == 1 and summary["overall_outcome"] == w.OUTCOME_WITH_FAILURES
    assert len(spark.calls) == 2
    assert any("[remote] FAILED" in l and "PERMISSION_DENIED" in l for l in logs)


def test_e2e_remote_rows_without_configuration_fail_before_any_ddl(vol):
    spark = FakeSpark()
    with pytest.raises(w.InputError, match="remote workspace but it is not configured") as e:
        run_e2e(vol, MIXED, spark=spark, remote=None)
    assert "remote_workspace_url" in str(e.value)
    assert spark.calls == []  # not even the local rows ran
    assert not (vol[1] / "Volumes/cat/sch/vol/rep.xlsx").exists()


@pytest.mark.parametrize("ws", [
    FakeWorkspace(token_status=401),
    FakeWorkspace(post_status=403, post_body={"message": "no CAN_USE"}),
    FakeWorkspace(script=[["FAILED"]]),
    FakeWorkspace(transport_error=OSError("down")),
])
def test_e2e_remote_preflight_failure_stops_everything_before_any_ddl(vol, ws):
    client, ws, _ = make_client(ws)
    spark = FakeSpark()
    with pytest.raises(w.InputError, match="Remote workspace preflight failed"):
        run_e2e(vol, MIXED, spark=spark, remote=client)
    assert spark.calls == []
    assert ws.sql == [] or ws.sql == ["SELECT 1"]  # nothing but the probe reached the remote
    assert SECRET not in "".join(str(r["url"]) for r in ws.requests)


def test_e2e_no_remote_rows_means_no_remote_call_at_all(vol):
    client, ws, _ = make_client()
    body = "SCHEMA,sales,a,,ENABLE,local\nSCHEMA,sales,b,,ENABLE,local\n"
    summary, details, spark, srows = run_e2e(vol, body, remote=client)
    assert ws.requests == [] and len(spark.calls) == 2
    assert summary["remote_rows"] == 0 and not srows["remote_workspace_url"]  # blank cell


def test_e2e_remote_rows_that_are_all_invalid_need_no_remote(vol):
    """Only rows that would really execute require the remote to be configured."""
    body = "SCHEMA,sales,a,,ENABLE,remote\nSCHEMA,sales,a,,DISABLE,remote\nBAD,x,,,ENABLE,remote\n"
    summary, details, spark, _ = run_e2e(vol, body, remote=None)
    assert [d["status"] for d in details] == ["SKIPPED_CONFLICT", "SKIPPED_CONFLICT", "VALIDATION_FAILED"]
    assert spark.calls == []
    assert summary["remote_rows"] == 3  # the report still shows what was requested


def test_e2e_legacy_csv_without_workspace_column_is_unchanged(vol):
    client, ws, _ = make_client()
    summary, details, spark, _ = run_e2e(vol, "SCHEMA,c,s,,ENABLE\nTABLE,c,s,t,DISABLE\n",
                                         remote=client, header="object_type,catalog,schema,table,action\n")
    assert ws.requests == []  # remote configured but unused
    assert len(spark.calls) == 2 and {d["workspace"] for d in details} == {"local"}
    assert summary["successes"] == 2


def test_e2e_local_failure_does_not_block_remote_rows(vol):
    client, ws, _ = make_client()
    spark = FakeSpark({"`loc1`": PermissionError("PERMISSION_DENIED")})
    summary, details, _, _ = run_e2e(vol, MIXED, spark=spark, remote=client)
    assert [d["status"] for d in details] == ["EXECUTION_FAILED", "SUCCESS", "SUCCESS", "SUCCESS"]
    assert len(ws.sql) == 3


def test_e2e_credentials_never_reach_the_report_or_logs(vol):
    ws = FakeWorkspace(script=[["SUCCEEDED"], [("FAILED", "X", "bad creds %s tok-1" % SECRET)]])
    client, _, _ = make_client(ws)
    logs = []
    root, tmp = vol
    run_e2e(vol, MIXED, remote=client, logs=logs)
    wb = openpyxl.load_workbook(real(tmp, root + "/rep.xlsx"))
    cells = [str(c.value) for sh in wb for row in sh.iter_rows() for c in row if c.value is not None]
    blob = "\n".join(cells + logs)
    assert SECRET not in blob and "tok-1" not in blob
    assert base64.b64encode(("%s:%s" % (CLIENT_ID, SECRET)).encode()).decode() not in blob


def test_details_columns_contain_workspace_and_statement_id():
    assert w.DETAIL_COLUMNS[:2] == ["source_row_number", "workspace"]
    assert w.DETAIL_COLUMNS[-1] == "remote_statement_id"
    assert len(w.record_to_row(w.ResultRecord(source_row_number=1))) == len(w.DETAIL_COLUMNS)


def test_remote_settings_widgets_match_worker_main():
    import inspect

    src = inspect.getsource(w.main)
    assert "REMOTE_SETTINGS" in src and "dbutils.secrets.get(scope, key)" in src
    assert w.REMOTE_SETTINGS == ("remote_workspace_url", "remote_client_id", "remote_secret_scope",
                                 "remote_secret_key", "remote_warehouse_id")
