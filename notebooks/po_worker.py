# Databricks notebook source
# MAGIC %md
# MAGIC # PO Control - Worker
# MAGIC Reads a CSV of Predictive Optimization (PO) operations from a Unity Catalog Volume,
# MAGIC validates it, executes one `ALTER ... {ENABLE|DISABLE} PREDICTIVE OPTIMIZATION` per
# MAGIC unique valid target (each in its own exception boundary) and writes an XLSX report
# MAGIC (`Summary` + `Details` sheets) to a Volume.
# MAGIC
# MAGIC **Task parameters (widgets):** `csv_path`, `report_path`, `volume_root`,
# MAGIC `overwrite_report`, `compute_type` (informational), `job_id`, `run_id`, and the optional remote
# MAGIC settings `remote_workspace_url`, `remote_client_id`, `remote_secret_scope`, `remote_secret_key`,
# MAGIC `remote_warehouse_id`.
# MAGIC
# MAGIC **Where a row runs:** the optional CSV column `workspace` is `local` (this notebook's compute) or
# MAGIC `remote` (a SQL warehouse in another workspace, via the SQL Statement API with an OAuth M2M token;
# MAGIC the client secret is read from a Databricks secret scope, never passed as a parameter). A CSV
# MAGIC without that column is entirely local.
# MAGIC
# MAGIC Failure policy: row-level problems are recorded in the report and do **not** fail the
# MAGIC task. Bootstrap problems (bad paths / CSV structure) and report-writing problems **do**
# MAGIC fail the task. DDL is never executed before paths and CSV structure are validated, the
# MAGIC report location is proven writable and, if any row is remote, the remote workspace has been
# MAGIC reached (token + `SELECT 1`).

# COMMAND ----------

import base64
import csv
import json
import os
import posixpath
import re
import shutil
import tempfile
import time
import traceback
import urllib.error
import urllib.parse
import urllib.request
import uuid
from collections import OrderedDict
from dataclasses import dataclass, field
from datetime import datetime, timezone
from typing import Callable, Dict, List, Optional, Tuple

# COMMAND ----------

# MAGIC %md
# MAGIC ## Constants and allowlists

# COMMAND ----------

REQUIRED_HEADER = ["object_type", "catalog", "schema", "table", "action"]
WORKSPACE_COLUMN = "workspace"  # optional; absent column = every row is local
CSV_COLUMNS = REQUIRED_HEADER + [WORKSPACE_COLUMN]
WORKSPACE_LOCAL = "local"
WORKSPACE_REMOTE = "remote"
ALLOWED_WORKSPACES = (WORKSPACE_LOCAL, WORKSPACE_REMOTE)

# object_type -> (identifier fields that must be present, fields that must be blank)
SCOPE_RULES = {
    "CATALOG": (("catalog",), ("schema", "table")),
    "SCHEMA": (("catalog", "schema"), ("table",)),
    "TABLE": (("catalog", "schema", "table"), ()),
}
ALLOWED_ACTIONS = ("ENABLE", "DISABLE")

# Allowlisted SQL templates. Only quoted identifiers and the allowlisted action are
# ever substituted; no CSV text is ever used as a SQL fragment.
SQL_TEMPLATES = {
    "CATALOG": "ALTER CATALOG {target} {action} PREDICTIVE OPTIMIZATION",
    "SCHEMA": "ALTER SCHEMA {target} {action} PREDICTIVE OPTIMIZATION",
    "TABLE": "ALTER TABLE {target} {action} PREDICTIVE OPTIMIZATION",
}

STATUS_SUCCESS = "SUCCESS"
STATUS_VALIDATION_FAILED = "VALIDATION_FAILED"
STATUS_EXECUTION_FAILED = "EXECUTION_FAILED"
STATUS_SKIPPED_DUPLICATE = "SKIPPED_DUPLICATE"
STATUS_SKIPPED_CONFLICT = "SKIPPED_CONFLICT"
ALL_STATUSES = (
    STATUS_SUCCESS,
    STATUS_VALIDATION_FAILED,
    STATUS_EXECUTION_FAILED,
    STATUS_SKIPPED_DUPLICATE,
    STATUS_SKIPPED_CONFLICT,
)

OUTCOME_ALL_SUCCEEDED = "ALL_SUCCEEDED"
OUTCOME_WITH_FAILURES = "COMPLETED_WITH_FAILURES"
OUTCOME_NO_ROWS = "NO_INPUT_ROWS"

DETAIL_COLUMNS = [
    "source_row_number",
    "workspace",
    "object_type",
    "catalog",
    "schema",
    "table",
    "fully_qualified_target",
    "requested_action",
    "generated_sql",
    "status",
    "start_timestamp",
    "end_timestamp",
    "error_class",
    "error_message",
    "reference_row_numbers",
    "remote_statement_id",
]

MAX_IDENTIFIER_LEN = 255
XLSX_MAX_CELL_CHARS = 32000  # Excel hard limit is 32767
TRUNCATION_MARKER = " ...[TRUNCATED - full text in task output]"
_ILLEGAL_XML_CHARS = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")
_FORMULA_PREFIXES = ("=", "+", "-", "@", "\t", "\r")


class InputError(Exception):
    """Bootstrap / file-level problem: no DDL may run, the task must fail."""


class ReportError(Exception):
    """Report generation or publishing problem: the task must fail."""


@dataclass
class ResultRecord:
    source_row_number: int
    workspace: str = ""
    remote_statement_id: str = ""
    object_type: str = ""
    catalog: str = ""
    schema: str = ""
    table: str = ""
    action: str = ""
    fully_qualified_target: str = ""
    generated_sql: str = ""
    status: str = ""
    start_timestamp: str = ""
    end_timestamp: str = ""
    error_class: str = ""
    error_message: str = ""
    references: List[int] = field(default_factory=list)
    target_key: Optional[Tuple] = None


# COMMAND ----------

# MAGIC %md
# MAGIC ## Time helper

# COMMAND ----------


def utc_now_iso() -> str:
    """Timezone-aware ISO-8601 timestamp (UTC)."""
    return datetime.now(timezone.utc).isoformat(timespec="milliseconds")


# COMMAND ----------

# MAGIC %md
# MAGIC ## Path validation (Volume-root boundary)

# COMMAND ----------


def _clean_volume_path(path: str, label: str) -> str:
    """Normalize a /Volumes path, rejecting traversal and malformed values."""
    if path is None or not str(path).strip():
        raise InputError("%s must not be empty" % label)
    raw = str(path).strip()
    if "\x00" in raw or "\n" in raw or "\r" in raw:
        raise InputError("%s contains illegal characters" % label)
    if not raw.startswith("/Volumes/"):
        raise InputError("%s must be an absolute /Volumes/... path: %r" % (label, raw))
    if ".." in raw.split("/"):
        raise InputError("%s must not contain '..' path segments: %r" % (label, raw))
    norm = posixpath.normpath(raw)
    # normpath keeps a leading '//' untouched; we already required '/Volumes/'.
    return norm


def validate_volume_root(volume_root: str) -> str:
    root = _clean_volume_path(volume_root, "volume_root")
    parts = [p for p in root.split("/") if p]
    # /Volumes/<catalog>/<schema>/<volume>[/subdir...]
    if len(parts) < 4:
        raise InputError(
            "volume_root must be at least /Volumes/<catalog>/<schema>/<volume>: %r" % root
        )
    return root


def validate_path_in_root(path: str, volume_root: str, label: str) -> str:
    root = validate_volume_root(volume_root)
    norm = _clean_volume_path(path, label)
    if not norm.startswith(root + "/"):
        raise InputError("%s %r is outside volume_root %r" % (label, norm, root))
    return norm


def validate_paths(csv_path: str, report_path: str, volume_root: str) -> Tuple[str, str]:
    """Validate both paths; returns normalized (csv_path, report_path)."""
    csv_norm = validate_path_in_root(csv_path, volume_root, "csv_path")
    rep_norm = validate_path_in_root(report_path, volume_root, "report_path")
    if not rep_norm.lower().endswith(".xlsx"):
        raise InputError("report_path must end with .xlsx: %r" % rep_norm)
    if not csv_norm.lower().endswith(".csv"):
        raise InputError("csv_path must end with .csv: %r" % csv_norm)
    if csv_norm == rep_norm or csv_norm.lower() == rep_norm.lower():
        raise InputError("report_path must differ from csv_path")
    if os.path.exists(csv_norm) and os.path.exists(rep_norm):
        try:
            if os.path.samefile(csv_norm, rep_norm):
                raise InputError("report_path resolves to the same file as csv_path")
        except OSError:
            pass
    return csv_norm, rep_norm


# COMMAND ----------

# MAGIC %md
# MAGIC ## SQL generation

# COMMAND ----------


def quote_identifier(name: str) -> str:
    """Backtick-quote a single identifier; embedded backticks are doubled."""
    return "`" + name.replace("`", "``") + "`"


def fully_qualified_target(object_type: str, catalog: str, schema: str, table: str) -> str:
    parts = {"CATALOG": [catalog], "SCHEMA": [catalog, schema], "TABLE": [catalog, schema, table]}[
        object_type
    ]
    return ".".join(quote_identifier(p) for p in parts)


def build_sql(object_type: str, action: str, catalog: str, schema: str, table: str) -> str:
    """Build the DDL only from allowlisted object_type/action and quoted identifiers."""
    if object_type not in SQL_TEMPLATES:
        raise ValueError("object_type not allowed: %r" % object_type)
    if action not in ALLOWED_ACTIONS:
        raise ValueError("action not allowed: %r" % action)
    target = fully_qualified_target(object_type, catalog, schema, table)
    return SQL_TEMPLATES[object_type].format(target=target, action=action)


# COMMAND ----------

# MAGIC %md
# MAGIC ## CSV reading (file-level validation)

# COMMAND ----------


def read_csv_records(csv_path: str) -> List[Tuple[int, List[str]]]:
    """Read the CSV and return [(source_row_number, fields)] for non-blank data rows.

    The header is row 1. Raises InputError for any file-level problem (missing/unreadable
    file, empty file, header not exactly the required column set). Every well-formed row is
    returned in the canonical order of CSV_COLUMNS (6 fields: a file without the optional
    `workspace` column gets "local" appended).
    """
    if not os.path.isfile(csv_path):
        raise InputError("CSV not found or not a file: %s" % csv_path)
    records = []  # type: List[Tuple[int, List[str]]]
    try:
        with open(csv_path, "r", encoding="utf-8-sig", newline="") as fh:
            reader = csv.reader(fh)
            header = None
            for idx, rec in enumerate(reader, start=1):
                if idx == 1:
                    header = [h.strip() for h in rec]
                    continue
                if not rec or all(c.strip() == "" for c in rec):
                    continue  # blank line: ignored, not counted
                records.append((idx, rec))
    except (OSError, UnicodeDecodeError, csv.Error) as exc:
        raise InputError("Cannot read CSV %s: %s: %s" % (csv_path, type(exc).__name__, exc))
    if header is None:
        raise InputError("CSV is empty (no header row): %s" % csv_path)
    _check_header(header)
    # Re-map well-formed rows to canonical column order (header may be in any order).
    # A row whose field count differs from the header is returned as MalformedFields so that
    # validate_row flags it (never silently treated as a short but valid row).
    has_workspace = WORKSPACE_COLUMN in header
    columns = CSV_COLUMNS if has_workspace else REQUIRED_HEADER
    pos = {name: header.index(name) for name in columns}
    mapped = []
    for idx, rec in records:
        if len(rec) == len(header):
            fields = [rec[pos[name]] for name in columns]
            if not has_workspace:
                fields.append(WORKSPACE_LOCAL)
            mapped.append((idx, fields))
        else:
            mapped.append((idx, MalformedFields(rec, expected=len(header))))
    return mapped


class MalformedFields(list):
    """A CSV row whose field count differs from the header's."""

    def __init__(self, fields, expected: int):
        super().__init__(fields)
        self.expected = expected


def _check_header(header: List[str]) -> None:
    if len(set(header)) != len(header):
        raise InputError("CSV header has duplicate columns: %s" % header)
    missing = [c for c in REQUIRED_HEADER if c not in header]
    extra = [c for c in header if c not in CSV_COLUMNS]
    if missing or extra:
        raise InputError(
            "CSV header must be exactly %s (the %s column is optional; without it every row is "
            "local); missing=%s extra=%s (found %s)"
            % (",".join(CSV_COLUMNS), WORKSPACE_COLUMN, missing, extra, header)
        )


# COMMAND ----------

# MAGIC %md
# MAGIC ## Row validation, duplicate and conflict policy

# COMMAND ----------


def _identifier_problem(label: str, value: str) -> Optional[str]:
    if len(value) > MAX_IDENTIFIER_LEN:
        return "%s is longer than %d characters" % (label, MAX_IDENTIFIER_LEN)
    if "." in value:
        return "%s must not contain '.'; put each name part in its own column" % label
    if "/" in value:
        return "%s must not contain '/'" % label
    if any(ch.isspace() or ord(ch) < 32 or ord(ch) == 127 for ch in value):
        return "%s must not contain whitespace or control characters" % label
    return None


def validate_row(source_row_number: int, fields: List[str]) -> ResultRecord:
    """Validate a single row. Returns a ResultRecord (status VALIDATION_FAILED or '')."""
    rec = ResultRecord(source_row_number=source_row_number)
    expected = getattr(fields, "expected", None)  # set by read_csv_records for header mismatches
    if expected is not None or len(fields) not in (len(REQUIRED_HEADER), len(CSV_COLUMNS)):
        rec.status = STATUS_VALIDATION_FAILED
        rec.error_class = "MalformedRow"
        rec.error_message = "Row has %d fields; expected %s" % (
            len(fields),
            expected if expected is not None else "%d or %d" % (len(REQUIRED_HEADER), len(CSV_COLUMNS)),
        )
        return rec

    # 5 fields = legacy row without the workspace column: implicitly local.
    vals = dict(zip(CSV_COLUMNS, [f.strip() for f in fields] + [WORKSPACE_LOCAL] * (len(CSV_COLUMNS) - len(fields))))
    rec.object_type = vals["object_type"].upper()
    rec.catalog, rec.schema, rec.table = vals["catalog"], vals["schema"], vals["table"]
    rec.action = vals["action"].upper()
    rec.workspace = vals[WORKSPACE_COLUMN].lower()

    problems = []
    if rec.workspace not in ALLOWED_WORKSPACES:
        problems.append(
            "workspace %r not allowed (expected %s)"
            % (vals[WORKSPACE_COLUMN], " or ".join(ALLOWED_WORKSPACES))
        )
    rules = SCOPE_RULES.get(rec.object_type)
    if rules is None:
        problems.append(
            "object_type %r not allowed (expected one of %s)"
            % (vals["object_type"], ", ".join(SCOPE_RULES))
        )
    if rec.action not in ALLOWED_ACTIONS:
        problems.append(
            "action %r not allowed (expected one of %s)"
            % (vals["action"], ", ".join(ALLOWED_ACTIONS))
        )
    if rules is not None:
        required, must_be_blank = rules
        for col in required:
            if not vals[col]:
                problems.append("%s is required for object_type %s" % (col, rec.object_type))
            else:
                p = _identifier_problem(col, vals[col])
                if p:
                    problems.append(p)
        for col in must_be_blank:
            if vals[col]:
                problems.append("%s must be blank for object_type %s" % (col, rec.object_type))

    if problems:
        rec.status = STATUS_VALIDATION_FAILED
        rec.error_class = "ValidationError"
        rec.error_message = "; ".join(problems)
        return rec

    # Valid: pre-compute target/key/sql. SQL is only *generated* here, never executed.
    rec.fully_qualified_target = fully_qualified_target(
        rec.object_type, rec.catalog, rec.schema, rec.table
    )
    rec.generated_sql = build_sql(rec.object_type, rec.action, rec.catalog, rec.schema, rec.table)
    # Unity Catalog identifiers are case-insensitive: normalize for duplicate detection.
    rec.target_key = (
        rec.workspace,  # the same name in the local and the remote workspace are different targets
        rec.object_type,
        rec.catalog.casefold(),
        rec.schema.casefold(),
        rec.table.casefold(),
    )
    return rec


def plan_records(raw_rows: List[Tuple[int, List[str]]]) -> Tuple[List[ResultRecord], List[ResultRecord]]:
    """Validate all rows and apply duplicate/conflict policy.

    Returns (all_records_in_row_order, executable_records_in_csv_order).
    Executable = the first row of each valid, non-conflicting target.
    """
    records = [validate_row(n, f) for n, f in raw_rows]

    groups = OrderedDict()  # type: OrderedDict
    for rec in records:
        if rec.status == STATUS_VALIDATION_FAILED:
            continue
        groups.setdefault(rec.target_key, []).append(rec)

    executable = []
    for key, members in groups.items():
        actions = {m.action for m in members}
        if len(actions) > 1:
            rows = [m.source_row_number for m in members]
            for m in members:
                m.status = STATUS_SKIPPED_CONFLICT
                m.error_class = "ConflictingActions"
                m.error_message = "Conflicting actions %s requested for the same target" % sorted(
                    actions
                )
                m.references = [r for r in rows if r != m.source_row_number]
                m.generated_sql = ""  # nothing will be executed for this target
            continue
        first = members[0]
        executable.append(first)
        for dup in members[1:]:
            dup.status = STATUS_SKIPPED_DUPLICATE
            dup.error_class = "DuplicateTarget"
            dup.error_message = "Identical request already handled at row %d" % first.source_row_number
            dup.references = [first.source_row_number]
    return records, executable


# COMMAND ----------

# MAGIC %md
# MAGIC ## Remote workspace (OAuth M2M + SQL Statement API)
# MAGIC Rows marked `remote` are sent as ordinary SQL statements to a SQL warehouse in another workspace
# MAGIC and waited for synchronously. The client secret is read from a secret scope, held only in memory
# MAGIC and scrubbed from every message, log line and report cell. A statement is submitted exactly once.

# COMMAND ----------

# The OAuth client secret is sent to this host, so only genuine Databricks workspace hosts are accepted.
REMOTE_HOST_SUFFIXES = (
    ".azuredatabricks.net",
    ".cloud.databricks.com",
    ".gcp.databricks.com",
    ".databricks.azure.cn",
    ".databricks.azure.us",
    ".cloud.databricks.us",
)
_HOST_RE = re.compile(r"^[a-z0-9]([a-z0-9\-]*[a-z0-9])?(\.[a-z0-9]([a-z0-9\-]*[a-z0-9])?)+$")
_SECRET_NAME_RE = re.compile(r"^[A-Za-z0-9_.\-]{1,128}$")
_WAREHOUSE_ID_RE = re.compile(r"^[A-Za-z0-9]{8,64}$")
REMOTE_SETTINGS = (
    "remote_workspace_url", "remote_client_id", "remote_secret_scope", "remote_secret_key",
    "remote_warehouse_id",
)
REMOTE_STATEMENT_TIMEOUT_SECONDS = 900  # warehouse cold start + the statement; never cancels on timeout
REMOTE_HTTP_TIMEOUT_SECONDS = 60
REMOTE_POLL_INTERVAL_SECONDS = 3
REMOTE_MAX_POLL_ERRORS = 3  # consecutive failed status polls tolerated (polling is read-only)
REMOTE_TOKEN_EARLY_REFRESH_SECONDS = 120
_REMOTE_OK_STATES = ("SUCCEEDED", "CLOSED")  # CLOSED = succeeded, result no longer fetchable
_REMOTE_PENDING_STATES = ("PENDING", "RUNNING")


class RemoteError(Exception):
    """A problem talking to the remote workspace. Messages never contain credentials."""

    def __init__(self, message, error_class="RemoteError", statement_id=""):
        super().__init__(message)
        self.error_class = error_class
        self.statement_id = statement_id


def normalize_workspace_url(value) -> str:
    """'https://<host>' for a Databricks workspace URL, else InputError.

    Query string and fragment (e.g. the '?o=123' browsers show) are dropped; credentials, other
    schemes, paths, odd ports and non-Databricks hosts are rejected.
    """
    raw = str(value if value is not None else "").strip()
    if not raw:
        raise InputError("remote_workspace_url must not be empty")
    if any(ch.isspace() or ord(ch) < 32 or ord(ch) == 127 for ch in raw):
        raise InputError("remote_workspace_url contains illegal characters")
    try:
        parts = urllib.parse.urlsplit(raw)
        port = parts.port
    except ValueError:
        raise InputError("remote_workspace_url is not a valid URL")
    if parts.scheme.lower() != "https":
        raise InputError("remote_workspace_url must start with https://")
    if parts.username is not None or parts.password is not None or "@" in parts.netloc:
        raise InputError("remote_workspace_url must not contain credentials")
    if port not in (None, 443):
        raise InputError("remote_workspace_url must not specify a port other than 443")
    if parts.path not in ("", "/"):
        raise InputError("remote_workspace_url must be just the workspace URL (https://<host>), no path")
    host = (parts.hostname or "").lower()
    if not _HOST_RE.match(host) or not host.endswith(REMOTE_HOST_SUFFIXES):
        raise InputError(
            "remote_workspace_url host %r is not a Databricks workspace host (expected one ending in %s)"
            % (host, ", ".join(REMOTE_HOST_SUFFIXES))
        )
    return "https://" + host


def _default_http(method, url, headers, body, timeout):
    """(status, body bytes) for any HTTP response, including 4xx/5xx. Network errors raise."""
    req = urllib.request.Request(url, data=body, headers=headers, method=method)
    try:
        with urllib.request.urlopen(req, timeout=timeout) as resp:  # noqa: S310 - https + host allowlist
            return resp.status, resp.read()
    except urllib.error.HTTPError as exc:
        return exc.code, exc.read()


class RemoteSqlClient:
    """Runs SQL on a warehouse in another workspace and waits for it (OAuth machine-to-machine).

    `secret_provider` returns the client secret and is only called when a token is first needed.
    `token_provider` (tests / diagnostics only) replaces the OAuth exchange with any bearer token.
    """

    def __init__(self, url, client_id, secret_provider, warehouse_id, http=None, sleep=time.sleep,
                 monotonic=time.monotonic, token_provider=None,
                 statement_timeout=REMOTE_STATEMENT_TIMEOUT_SECONDS,
                 poll_interval=REMOTE_POLL_INTERVAL_SECONDS, log=print):
        self.url = normalize_workspace_url(url)
        self.client_id = str(client_id).strip()
        self.warehouse_id = str(warehouse_id).strip()
        self._secret_provider = secret_provider
        self._http = http or _default_http
        self._sleep, self._mono = sleep, monotonic
        self._token_provider = token_provider
        self._statement_timeout = statement_timeout
        self._poll_interval = poll_interval
        self._log = log
        self._secret = None
        self._token = None
        self._token_expiry = 0.0
        self._sensitive = []  # every credential-bearing string ever seen, scrubbed from messages

    # ---- helpers
    def _remember(self, value):
        if value and value not in self._sensitive:
            self._sensitive.append(value)

    def _redact(self, text) -> str:
        out = str(text)
        for v in sorted(self._sensitive, key=len, reverse=True):
            out = out.replace(v, "***")
        return out

    def _call(self, method, path, headers, body=None):
        """(status, parsed JSON or {'raw': text}). Transport failures become RemoteError."""
        try:
            status, raw = self._http(method, self.url + path, headers, body, REMOTE_HTTP_TIMEOUT_SECONDS)
        except Exception as exc:  # noqa: BLE001 - any transport problem
            raise RemoteError(
                "Remote workspace %s unreachable: %s: %s" % (self.url, type(exc).__name__, self._redact(exc)),
                "RemoteConnectionError",
            )
        text = raw.decode("utf-8", "replace") if isinstance(raw, (bytes, bytearray)) else str(raw)
        try:
            data = json.loads(text) if text.strip() else {}
        except ValueError:
            data = {"raw": text[:500]}
        return status, data if isinstance(data, dict) else {"raw": text[:500]}

    def _hint(self, status, data) -> str:
        blob = json.dumps(data).lower()
        if "ip acl" in blob or "ip access list" in blob or "ip allow" in blob:
            return (" The remote workspace's IP access list is blocking this compute's egress IP; "
                    "allow it there (serverless: its NCC/stable IPs).")
        return ""

    # ---- OAuth
    def _get_token(self, force_refresh=False) -> str:
        if self._token_provider:
            return self._token_provider()
        if self._token and not force_refresh and self._mono() < self._token_expiry - REMOTE_TOKEN_EARLY_REFRESH_SECONDS:
            return self._token
        if self._secret is None:
            try:
                secret = self._secret_provider()
            except Exception as exc:  # noqa: BLE001 - any secret-store problem
                raise RemoteError(
                    "Cannot read the OAuth client secret from the secret scope: %s: %s"
                    % (type(exc).__name__, exc),
                    "RemoteAuthError",
                )
            if not secret:
                raise RemoteError("The OAuth client secret is empty", "RemoteAuthError")
            self._secret = str(secret)
            self._remember(self._secret)
        basic = base64.b64encode(("%s:%s" % (self.client_id, self._secret)).encode("utf-8")).decode("ascii")
        self._remember(basic)
        status, data = self._call(
            "POST", "/oidc/v1/token",
            {"Authorization": "Basic " + basic, "Content-Type": "application/x-www-form-urlencoded",
             "Accept": "application/json"},
            b"grant_type=client_credentials&scope=all-apis",
        )
        token = data.get("access_token") if status == 200 else None
        if not token:
            detail = data.get("error_description") or data.get("error") or data.get("raw") or ""
            raise RemoteError(
                "OAuth token request to %s failed (HTTP %s): %s%s"
                % (self.url, status, self._redact(detail)[:300], self._hint(status, data)),
                "RemoteAuthError",
            )
        self._remember(token)
        self._token = token
        try:
            ttl = float(data.get("expires_in", 3600))
        except (TypeError, ValueError):
            ttl = 3600.0
        self._token_expiry = self._mono() + ttl
        return token

    def _api(self, method, path, body=None):
        """Authenticated call; one transparent re-authentication if the token was rejected (401)."""
        payload = json.dumps(body).encode("utf-8") if body is not None else None
        for attempt in (0, 1):
            token = self._get_token(force_refresh=attempt == 1)
            self._remember(token)
            headers = {"Authorization": "Bearer " + token, "Accept": "application/json"}
            if payload is not None:
                headers["Content-Type"] = "application/json"
            status, data = self._call(method, path, headers, payload)
            if status != 401 or attempt == 1 or self._token_provider:
                return status, data
        return status, data  # pragma: no cover

    # ---- statements
    @staticmethod
    def sql_condition(error_code: str, message: str) -> str:
        """The meaningful error condition of a failed remote statement.

        Found by live testing: the API reports a generic error_code (e.g. BAD_REQUEST) and puts the real
        condition at the start of the message, as '[TABLE_OR_VIEW_NOT_FOUND] ...' or 'PERMISSION_DENIED: ...'.
        """
        m = re.match(r"\s*\[([A-Z][A-Z0-9_.]*)\]", message or "") or re.match(r"\s*([A-Z][A-Z0-9_]{3,}):", message or "")
        return m.group(1) if m else (error_code or "FAILED")

    @staticmethod
    def _status_of(data):
        st = data.get("status") or {}
        return st.get("state", ""), st.get("error") or {}

    def execute(self, sql: str, statement_timeout=None) -> str:
        """Submit `sql` once, wait for it to finish; returns the statement_id on success.

        Raises RemoteError (with .statement_id when one exists) on failure. A timeout does NOT
        cancel the statement: its final state is unknown and is reported as such.
        """
        status, data = self._api("POST", "/api/2.0/sql/statements/", {
            "statement": sql, "warehouse_id": self.warehouse_id,
            "wait_timeout": "30s", "on_wait_timeout": "CONTINUE",
        })
        if not 200 <= status < 300:
            detail = data.get("message") or data.get("error_description") or data.get("raw") or ""
            raise RemoteError(
                "Remote statement submission failed (HTTP %s): %s%s"
                % (status, self._redact(detail)[:500], self._hint(status, data)),
                "RemoteApiError:HTTP_%s" % status,
            )
        sid = data.get("statement_id") or ""
        if not sid:
            raise RemoteError("Remote workspace returned no statement_id", "RemoteApiError")
        deadline = self._mono() + (self._statement_timeout if statement_timeout is None else statement_timeout)
        poll_errors = 0
        while True:
            state, err = self._status_of(data)
            if state in _REMOTE_OK_STATES:
                return sid
            if state == "FAILED":
                code = err.get("error_code") or "FAILED"
                msg = err.get("message", "statement failed")
                raise RemoteError("%s: %s" % (code, self._redact(msg)),
                                  "RemoteSqlError:%s" % self.sql_condition(code, msg), sid)
            if state == "CANCELED":
                raise RemoteError("Remote statement was canceled", "RemoteSqlError:CANCELED", sid)
            if state not in _REMOTE_PENDING_STATES:
                raise RemoteError("Remote statement is in unexpected state %r" % state, "RemoteApiError", sid)
            if self._mono() >= deadline:
                raise RemoteError(
                    "Timed out waiting for the remote statement (last state %s); it was NOT cancelled and "
                    "its final outcome is unknown. Check statement_id %s on %s" % (state, sid, self.url),
                    "RemoteTimeout", sid)
            self._sleep(self._poll_interval)
            try:
                st, nxt = self._api("GET", "/api/2.0/sql/statements/%s" % urllib.parse.quote(sid, safe=""))
                if not 200 <= st < 300:
                    raise RemoteError("HTTP %s while polling" % st, "RemoteConnectionError", sid)
            except RemoteError as exc:
                poll_errors += 1
                if poll_errors >= REMOTE_MAX_POLL_ERRORS:
                    raise RemoteError(
                        "Lost contact with the remote workspace while waiting (%s); the statement may still "
                        "be running, its final outcome is unknown. Check statement_id %s on %s"
                        % (exc, sid, self.url), "RemoteConnectionError", sid)
                self._log("remote poll for %s failed (%s); retrying" % (sid, exc))
                continue
            poll_errors = 0
            data = nxt

    def preflight(self) -> None:
        """Prove URL, credentials, warehouse id and CAN USE work, before any DDL runs anywhere."""
        self.execute("SELECT 1")


def build_remote_client(url, client_id, secret_scope, secret_key, warehouse_id, secret_getter, **kwargs):
    """RemoteSqlClient from the five remote_* settings; None if all are blank; InputError if partial/bad."""
    values = dict(zip(REMOTE_SETTINGS, (url, client_id, secret_scope, secret_key, warehouse_id)))
    values = {k: str(v if v is not None else "").strip() for k, v in values.items()}
    if not any(values.values()):
        return None
    missing = [k for k, v in values.items() if not v]
    if missing:
        raise InputError("remote workspace settings are incomplete, missing: %s" % ", ".join(missing))
    normalized = normalize_workspace_url(values["remote_workspace_url"])
    if any(ch.isspace() or ord(ch) < 32 for ch in values["remote_client_id"]) or len(values["remote_client_id"]) > 128:
        raise InputError("remote_client_id must be at most 128 characters without whitespace")
    for label in ("remote_secret_scope", "remote_secret_key"):
        if not _SECRET_NAME_RE.match(values[label]):
            raise InputError("%s must be 1-128 characters of letters, digits, '_', '-' or '.'" % label)
    if not _WAREHOUSE_ID_RE.match(values["remote_warehouse_id"]):
        raise InputError("remote_warehouse_id must be 8-64 letters/digits (the warehouse ID, not its name)")
    scope, key = values["remote_secret_scope"], values["remote_secret_key"]
    return RemoteSqlClient(
        normalized, values["remote_client_id"], lambda: secret_getter(scope, key),
        values["remote_warehouse_id"], **kwargs
    )


# COMMAND ----------

# MAGIC %md
# MAGIC ## Execution (one exception boundary per statement, no retries)

# COMMAND ----------


def _error_class_of(exc: BaseException) -> str:
    for attr in ("getCondition", "getErrorClass"):
        fn = getattr(exc, attr, None)
        if callable(fn):
            try:
                val = fn()
            except Exception:  # noqa: BLE001 - diagnostics only
                val = None
            if val:
                return "%s:%s" % (type(exc).__name__, val)
    return type(exc).__name__


def execute_records(
    executable: List[ResultRecord],
    sql_runner: Callable[[str], object],
    log: Callable[[str], None] = print,
    clock: Callable[[], str] = utc_now_iso,
    remote_runner: Optional[Callable[[str], str]] = None,
) -> None:
    """Execute each record's SQL once, in CSV order. A failure never stops later records.

    local rows go to `sql_runner` (this notebook's compute); remote rows go to `remote_runner`
    (the remote warehouse; it blocks until the statement finishes and returns its statement_id).
    """
    for rec in executable:
        remote = rec.workspace == WORKSPACE_REMOTE
        tag = "remote" if remote else "local"
        rec.start_timestamp = clock()
        try:
            if remote:
                if remote_runner is None:  # run_worker's preflight prevents this; never fall back to local
                    raise RemoteError("remote workspace is not configured", "RemoteNotConfigured")
                rec.remote_statement_id = remote_runner(rec.generated_sql) or ""  # exactly one attempt
            else:
                sql_runner(rec.generated_sql)  # exactly one attempt: no automatic retry
            rec.status = STATUS_SUCCESS
        except Exception as exc:  # noqa: BLE001 - row-level isolation is intentional
            rec.status = STATUS_EXECUTION_FAILED
            if isinstance(exc, RemoteError):
                rec.error_class = exc.error_class
                rec.remote_statement_id = exc.statement_id or rec.remote_statement_id
            else:
                rec.error_class = _error_class_of(exc)
            rec.error_message = str(exc) or type(exc).__name__
            log("Row %d [%s] FAILED: %s -> %s%s" % (
                rec.source_row_number, tag, rec.generated_sql, rec.error_class,
                " (statement_id=%s)" % rec.remote_statement_id if rec.remote_statement_id else ""))
        else:
            log("Row %d [%s] SUCCESS: %s%s" % (
                rec.source_row_number, tag, rec.generated_sql,
                " (statement_id=%s)" % rec.remote_statement_id if rec.remote_statement_id else ""))
        finally:
            rec.end_timestamp = clock()


# COMMAND ----------

# MAGIC %md
# MAGIC ## Summary

# COMMAND ----------


def count_statuses(records: List[ResultRecord]) -> Dict[str, int]:
    counts = {s: 0 for s in ALL_STATUSES}
    for r in records:
        counts[r.status] = counts.get(r.status, 0) + 1
    return counts


def build_summary(
    records: List[ResultRecord],
    csv_path: str,
    report_path: str,
    start_ts: str,
    end_ts: str,
    run_id: str = "",
    job_id: str = "",
    report_status: str = "WRITTEN",
    compute_type: str = "",
    remote_url: str = "",
    remote_warehouse_id: str = "",
) -> "OrderedDict[str, object]":
    c = count_statuses(records)
    remote = [r for r in records if r.workspace == WORKSPACE_REMOTE]
    attempted = c[STATUS_SUCCESS] + c[STATUS_EXECUTION_FAILED]
    problems = (
        c[STATUS_VALIDATION_FAILED]
        + c[STATUS_EXECUTION_FAILED]
        + c[STATUS_SKIPPED_CONFLICT]
    )
    if not records:
        outcome = OUTCOME_NO_ROWS
    elif problems:
        outcome = OUTCOME_WITH_FAILURES
    else:
        outcome = OUTCOME_ALL_SUCCEEDED
    s = OrderedDict()
    s["run_id"] = run_id
    s["job_id"] = job_id
    s["compute_type"] = compute_type
    s["input_csv_path"] = csv_path
    s["report_path"] = report_path
    s["remote_workspace_url"] = remote_url
    s["remote_warehouse_id"] = remote_warehouse_id
    s["start_timestamp"] = start_ts
    s["end_timestamp"] = end_ts
    s["total_input_rows"] = len(records)
    s["local_rows"] = sum(1 for r in records if r.workspace == WORKSPACE_LOCAL)
    s["remote_rows"] = len(remote)
    s["remote_successes"] = sum(1 for r in remote if r.status == STATUS_SUCCESS)
    s["remote_execution_failures"] = sum(1 for r in remote if r.status == STATUS_EXECUTION_FAILED)
    s["valid_unique_targets_attempted"] = attempted
    s["successes"] = c[STATUS_SUCCESS]
    s["validation_failures"] = c[STATUS_VALIDATION_FAILED]
    s["execution_failures"] = c[STATUS_EXECUTION_FAILED]
    s["duplicates_skipped"] = c[STATUS_SKIPPED_DUPLICATE]
    s["conflicts_skipped"] = c[STATUS_SKIPPED_CONFLICT]
    s["report_status"] = report_status
    s["overall_outcome"] = outcome
    s["note"] = (
        "SUCCESS means the ALTER statement returned without error; it does not prove the "
        "effective PO behavior of every child object. Parent settings do not override "
        "explicitly configured children."
    )
    return s


# COMMAND ----------

# MAGIC %md
# MAGIC ## XLSX report

# COMMAND ----------


def sanitize_cell(value, log: Callable[[str], None] = print, where: str = ""):
    """Make a value safe for XLSX: neutralize formulas, drop illegal chars, cap length."""
    if not isinstance(value, str):
        return value
    text = _ILLEGAL_XML_CHARS.sub("\ufffd", value)
    if text.startswith(_FORMULA_PREFIXES):
        text = "'" + text  # prevents interpretation as a formula
    if len(text) > XLSX_MAX_CELL_CHARS:
        log("TRUNCATED cell (%s), full text follows:\n%s" % (where, value))
        text = text[: XLSX_MAX_CELL_CHARS - len(TRUNCATION_MARKER)] + TRUNCATION_MARKER
    return text


def record_to_row(rec: ResultRecord) -> List[object]:
    return [
        rec.source_row_number,
        rec.workspace,
        rec.object_type,
        rec.catalog,
        rec.schema,
        rec.table,
        rec.fully_qualified_target,
        rec.action,
        rec.generated_sql,
        rec.status,
        rec.start_timestamp,
        rec.end_timestamp,
        rec.error_class,
        rec.error_message,
        ", ".join(str(r) for r in rec.references),
        rec.remote_statement_id,
    ]


def build_workbook(records: List[ResultRecord], summary, log: Callable[[str], None] = print):
    import openpyxl  # imported lazily so tests of pure logic do not need it

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Summary"
    ws.append(["metric", "value"])
    for k, v in summary.items():
        ws.append([k, sanitize_cell(v, log, "Summary/%s" % k)])
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 100

    wd = wb.create_sheet("Details")
    wd.append(DETAIL_COLUMNS)
    for rec in sorted(records, key=lambda r: r.source_row_number):
        row = record_to_row(rec)
        wd.append(
            [
                sanitize_cell(v, log, "Details/row %d/%s" % (rec.source_row_number, DETAIL_COLUMNS[i]))
                for i, v in enumerate(row)
            ]
        )
    wd.freeze_panes = "A2"
    widths = {
        "source_row_number": 10, "workspace": 11, "object_type": 12, "catalog": 16, "schema": 16,
        "table": 16, "fully_qualified_target": 44, "requested_action": 12, "generated_sql": 70,
        "status": 20, "start_timestamp": 28, "end_timestamp": 28, "error_class": 28,
        "error_message": 80, "reference_row_numbers": 16, "remote_statement_id": 40,
    }
    for idx, name in enumerate(DETAIL_COLUMNS, start=1):
        wd.column_dimensions[openpyxl.utils.get_column_letter(idx)].width = widths[name]
    return wb


def verify_report(path: str, expected_detail_rows: int) -> None:
    """Re-open the published workbook and check structure and row counts."""
    import openpyxl

    wb = openpyxl.load_workbook(path, read_only=True)
    try:
        if wb.sheetnames != ["Summary", "Details"]:
            raise ReportError("Unexpected sheets in report: %s" % wb.sheetnames)
        n = sum(1 for _ in wb["Details"].iter_rows(min_row=2, values_only=True))
        if n != expected_detail_rows:
            raise ReportError("Details rows %d != expected %d" % (n, expected_detail_rows))
    finally:
        wb.close()


def preflight_report_location(report_path: str, overwrite: bool) -> None:
    """Prove the report can be written BEFORE any DDL runs."""
    if os.path.exists(report_path) and not overwrite:
        raise InputError(
            "report_path already exists and overwrite_report is false: %s" % report_path
        )
    if os.path.isdir(report_path):
        raise InputError("report_path is a directory: %s" % report_path)
    rdir = posixpath.dirname(report_path)
    try:
        os.makedirs(rdir, exist_ok=True)
        probe = posixpath.join(rdir, ".po_probe_%s.tmp" % uuid.uuid4().hex)
        with open(probe, "wb") as fh:
            fh.write(b"probe")
        os.remove(probe)
    except OSError as exc:
        raise InputError(
            "Report directory %s is not writable: %s. Check the run identity's WRITE VOLUME "
            "privilege. If this runs on classic compute (job cluster / all-purpose) and ALL "
            "Unity Catalog access fails, the cluster's egress IP may be blocked by the workspace "
            "IP access list (look for '403 ... blocked by Databricks IP ACL' in the cluster logs)."
            % (rdir, exc)
        )
    try:
        import openpyxl  # noqa: F401
    except ImportError as exc:
        raise InputError("openpyxl is required to write the XLSX report: %s" % exc)


def write_report(records, summary, report_path: str, overwrite: bool, log=print) -> None:
    """Build the workbook locally, stage it next to the target, publish, then verify."""
    try:
        wb = build_workbook(records, summary, log)
        with tempfile.TemporaryDirectory() as tmpdir:
            local = os.path.join(tmpdir, "report.xlsx")
            wb.save(local)  # local disk supports random access; Volumes may not
            rdir = posixpath.dirname(report_path)
            staged = posixpath.join(rdir, ".%s.%s.tmp" % (posixpath.basename(report_path), uuid.uuid4().hex))
            shutil.copyfile(local, staged)
            try:
                if os.path.exists(report_path) and not overwrite:
                    raise ReportError("report_path appeared during the run: %s" % report_path)
                try:
                    os.replace(staged, report_path)
                except OSError as exc:
                    log("os.replace unsupported here (%s); falling back to copy" % exc)
                    shutil.copyfile(staged, report_path)
            finally:
                if os.path.exists(staged):
                    try:
                        os.remove(staged)
                    except OSError as exc:
                        log("Could not remove staging file %s: %s" % (staged, exc))
        verify_report(report_path, len(records))
    except ReportError:
        raise
    except Exception as exc:  # noqa: BLE001 - re-raised as ReportError after logging
        log("REPORT WRITE FAILED:\n" + traceback.format_exc())
        raise ReportError("Failed to write report %s: %s: %s" % (report_path, type(exc).__name__, exc))


# COMMAND ----------

# MAGIC %md
# MAGIC ## Orchestration

# COMMAND ----------


def parse_bool(value: str) -> bool:
    return str(value).strip().lower() in ("true", "1", "yes", "y")


def run_worker(
    csv_path: str,
    report_path: str,
    volume_root: str,
    sql_runner: Callable[[str], object],
    overwrite: bool = False,
    job_id: str = "",
    run_id: str = "",
    log: Callable[[str], None] = print,
    compute_type: str = "",
    remote_client: Optional[RemoteSqlClient] = None,
) -> "OrderedDict[str, object]":
    """Full worker flow. Raises InputError/ReportError for task-failing conditions."""
    start_ts = utc_now_iso()
    # ---- bootstrap: nothing below may execute DDL until it all passes ----
    csv_norm, report_norm = validate_paths(csv_path, report_path, volume_root)
    preflight_report_location(report_norm, overwrite)
    raw_rows = read_csv_records(csv_norm)
    records, executable = plan_records(raw_rows)
    n_remote = sum(1 for r in executable if r.workspace == WORKSPACE_REMOTE)
    log(
        "Parsed %d rows; %d valid unique targets to execute (%d local, %d remote)"
        % (len(records), len(executable), len(executable) - n_remote, n_remote)
    )
    if n_remote:  # still before any DDL: a bad remote must not leave a half-applied batch
        if remote_client is None:
            raise InputError(
                "%d row(s) target the remote workspace but it is not configured; set %s"
                % (n_remote, ", ".join(REMOTE_SETTINGS))
            )
        try:
            remote_client.preflight()
        except RemoteError as exc:
            raise InputError("Remote workspace preflight failed (%s): %s" % (exc.error_class, exc))
        log("Remote workspace reachable: %s (warehouse %s)" % (remote_client.url, remote_client.warehouse_id))
    # ---- execution ----
    execute_records(executable, sql_runner, log, remote_runner=remote_client.execute if n_remote else None)
    # ---- report ----
    summary = build_summary(
        records, csv_norm, report_norm, start_ts, utc_now_iso(), run_id=run_id, job_id=job_id,
        compute_type=compute_type,
        remote_url=remote_client.url if n_remote else "",
        remote_warehouse_id=remote_client.warehouse_id if n_remote else "",
    )
    write_report(records, summary, report_norm, overwrite, log)
    return summary


def main():  # pragma: no cover - runs only inside Databricks
    dbutils.widgets.text("csv_path", "")
    dbutils.widgets.text("report_path", "")
    dbutils.widgets.text("volume_root", "")
    dbutils.widgets.text("overwrite_report", "false")
    dbutils.widgets.text("job_id", "")
    dbutils.widgets.text("run_id", "")
    dbutils.widgets.text("compute_type", "")
    for name in REMOTE_SETTINGS:
        dbutils.widgets.text(name, "")
    get = dbutils.widgets.get

    try:
        remote_client = build_remote_client(
            *(get(name) for name in REMOTE_SETTINGS),
            secret_getter=lambda scope, key: dbutils.secrets.get(scope, key),
        )
        summary = run_worker(
            csv_path=get("csv_path"),
            report_path=get("report_path"),
            volume_root=get("volume_root"),
            sql_runner=lambda sql: spark.sql(sql),  # noqa: F821 - provided by the runtime
            overwrite=parse_bool(get("overwrite_report")),
            job_id=get("job_id"),
            run_id=get("run_id"),
            compute_type=get("compute_type"),
            remote_client=remote_client,
        )
    except (InputError, ReportError) as exc:
        print("PO WORKER FAILED (%s): %s" % (type(exc).__name__, exc))
        raise  # surfaces as a failed task

    print("=" * 70)
    for k, v in summary.items():
        if k != "note":
            print("%-34s %s" % (k, v))
    print("Report written to: %s" % summary["report_path"])
    dbutils.notebook.exit(json.dumps(summary))


# COMMAND ----------

if __name__ == "__main__":
    main()
