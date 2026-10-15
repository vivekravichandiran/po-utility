#!/usr/bin/env python3
"""Integration driver: behaves like a person using the wrapper notebook.

1. uploads a scenario CSV to the Volume, 2. runs the po_wrapper notebook (via a one-time run
with widget values), 3. prints wrapper result, 4. downloads and prints the XLSX report.

Usage: integration/it_run.py <profile> <csv_file> [--csv-name NAME.csv]
       [--compute serverless|job_cluster|all_purpose] [--cluster-id ID] [--label TEXT]
       [--set widget=value]

The wrapper has four widgets only: csv_path, report_path, compute_type, all_purpose_cluster_id.
report_path is the Volume folder; the wrapper names the report po_report_<timestamp>.xlsx.
"""
import argparse
import json
import os
import subprocess
import sys
import tempfile
import time
import warnings

warnings.filterwarnings("ignore")
from databricks.sdk import WorkspaceClient  # noqa: E402

ROOT = "/Volumes/po_ctl_admin/ops/po_control"


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("profile")
    ap.add_argument("csv_file")
    ap.add_argument("--csv-name", default=None)
    ap.add_argument("--compute", default=None, help="serverless|job_cluster|all_purpose (sets compute_type widget)")
    ap.add_argument("--cluster-id", default=None, help="all_purpose_cluster_id widget")
    ap.add_argument("--label", default="", help="free-text label for the run log")
    ap.add_argument("--set", action="append", default=[], help="override a wrapper widget key=value")
    a = ap.parse_args()

    w = WorkspaceClient(profile=a.profile)
    csv_name = a.csv_name or os.path.basename(a.csv_file)
    csv_vol = "%s/%s" % (ROOT, csv_name)
    if a.csv_file != "-":  # "-" = do not upload (e.g. to test a missing CSV)
        subprocess.run(["databricks", "fs", "cp", a.csv_file, "dbfs:" + csv_vol, "--overwrite",
                        "--profile", a.profile], check=True, capture_output=True)
    params = {
        "csv_path": csv_vol,
        "report_path": ROOT,  # a folder; the wrapper generates the file name
    }
    if a.compute:
        params["compute_type"] = a.compute
    if a.cluster_id:
        params["all_purpose_cluster_id"] = a.cluster_id
    for kv in a.set:
        k, v = kv.split("=", 1)
        params[k] = v
    print("wrapper widgets:", json.dumps(params))
    body = {
        "run_name": "po_wrapper (it_run)",
        "tasks": [{"task_key": "wrapper", "notebook_task": {
            "notebook_path": "/Shared/po_control/po_wrapper", "source": "WORKSPACE",
            "base_parameters": params}}],
    }
    resp = w.api_client.do("POST", "/api/2.1/jobs/runs/submit", body=body)
    rid = resp["run_id"]
    print("wrapper run_id:", rid)
    while True:
        r = w.api_client.do("GET", "/api/2.1/jobs/runs/get", query={"run_id": rid})
        st = r["state"]
        if st["life_cycle_state"] in ("TERMINATED", "SKIPPED", "INTERNAL_ERROR"):
            break
        time.sleep(8)
    wrapper_url = r.get("run_page_url", "")
    print("wrapper state:", st["life_cycle_state"], st.get("result_state"), st.get("state_message", "")[:300])
    task_run = r["tasks"][0]["run_id"]
    out = w.api_client.do("GET", "/api/2.1/jobs/runs/get-output", query={"run_id": task_run})
    if out.get("error"):
        print("wrapper ERROR:", out["error"])
        print((out.get("error_trace") or "")[-1500:])
    res = out.get("notebook_output", {}).get("result")
    if res:
        res = json.loads(res)
        print("wrapper result:", json.dumps(res, indent=1))
        if res.get("worker_summary"):
            print("worker summary:", json.dumps(json.loads(res["worker_summary"]), indent=1))
    report_file = (res or {}).get("report_path")
    if st.get("result_state") == "SUCCESS" and report_file:
        show_report(a.profile, report_file)
    log_run(a, params, rid, wrapper_url, st, res if res else None)
    print("CSV:", csv_vol, "| REPORT:", report_file)


def log_run(a, params, rid, wrapper_url, st, res):
    rec = {
        "label": a.label, "compute_type": params.get("compute_type", "serverless"),
        "csv": os.path.basename(a.csv_file),
        "wrapper_run_id": rid, "wrapper_run_url": wrapper_url,
        "wrapper_state": "%s/%s" % (st["life_cycle_state"], st.get("result_state")),
        "job_id": (res or {}).get("job_id"), "job_name": (res or {}).get("job_name"),
        "job_action": (res or {}).get("job_action"),
        "worker_run_id": (res or {}).get("run_id"), "worker_run_url": (res or {}).get("run_page_url"),
        "worker_state": ((res or {}).get("state"), (res or {}).get("result_state")),
        "report": (res or {}).get("report_path"),
    }
    with open(os.path.join(os.path.dirname(os.path.abspath(__file__)), "runs_log.jsonl"), "a") as fh:
        fh.write(json.dumps(rec) + "\n")


def show_report(profile, vol_path):
    import openpyxl

    with tempfile.TemporaryDirectory() as d:
        local = os.path.join(d, "r.xlsx")
        p = subprocess.run(["databricks", "fs", "cp", "dbfs:" + vol_path, local, "--profile", profile],
                           capture_output=True, text=True)
        if p.returncode:
            print("could not download report:", p.stderr[-300:])
            return
        wb = openpyxl.load_workbook(local)
        print("--- Summary")
        for r in wb["Summary"].iter_rows(min_row=2, values_only=True):
            if r[0] != "note":
                print("  %-34s %s" % r)
        print("--- Details")
        for r in wb["Details"].iter_rows(min_row=1, values_only=True):
            print("  ", r[0], r[1], r[5], r[6], "|", r[8], "|", (r[11] or ""), "|", (r[12] or "")[:90], "|", r[13])


if __name__ == "__main__":
    sys.exit(main())
